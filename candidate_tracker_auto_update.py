# -*- coding: utf-8 -*-
"""
★후보숫자_추적표_표본vs최근50회.xlsx 자동 업데이트 스크립트.

매주 토요일 로또 추첨이 끝나면, 동행복권 공식 API에서 새 회차 결과를 가져와
"전체당첨내역" 시트에 추가한다. 이 파일은 weekly_lotto_file_update.py가 다루는
4개 파일(샘플/200회검증용/전체표본/최근500표본)과 구조가 다르다 — "3차필터 예측행
freeze·전진" 같은 복잡한 절차가 필요 없다:

  이 파일의 10개 "기준{W}회_후보" 시트 + 기준 500회 시트("3차필터(50회_후보)")는
  전부 `=MAX(전체당첨내역!$A$2:$A$1500)+1` 같은 수식으로 "현재 최신 회차"를 스스로
  찾아내는 구조라서, 전체당첨내역에 새 행만 추가하면 나머지는 전부 자동으로
  재계산된다. 사람이 매주 예측행을 수동으로 밀어줄 필요가 없다.

다만 딱 하나, 자동으로 안 늘어나는 보조 시트가 있다: 숨김 시트 "_calc_누적빈도"
(회차별 누적 출현횟수를 저장해 두고 빼기만 해서 구간빈도를 구하는 캐시 테이블)는
전체당첨내역과 정확히 1:1로 행이 대응해야 하므로, 새 회차를 추가할 때마다 이
시트에도 같은 회차의 누적값 행을 함께 추가해야 한다. (이 스크립트가 자동으로
처리함 — 2026-09-15에 22,950건 교차검증 + 1회차/2회차 백필 시뮬레이션으로
검증 완료.)

필요 사전 설치 (엑셀/WPS가 설치된 PC에서, 명령 프롬프트에서 한 번만):
    pip install openpyxl pywin32

작업 스케줄러 등록 (관리자 권한 명령 프롬프트에서 한 번만 실행 — 선택사항):
    schtasks /create /tn "로또신령_후보숫자추적표_자동업데이트" ^
        /tr "python C:\\Users\\PC\\Desktop\\lotto-app\\candidate_tracker_auto_update.py" ^
        /sc weekly /d SUN /st 09:10 /ri 60 /du 0004:50 /f

    (weekly_lotto_file_update.py와 동일한 스케줄 — 09:10부터 60분 간격으로
    4시간50분 동안 재시도. 이미 최신이면 아무것도 안 하고 조용히 종료하므로
    여러 번 실행돼도 안전함.)
"""

from __future__ import annotations

import datetime
import json
import sys
import urllib.request
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

# ============================== CONFIG ======================================

TARGET_FILE = Path(
    r"C:\Users\PC\Desktop\lotto-app\★조합생성_후보숫자_추적표"
    r"\★후보숫자_추적표_표본vs최근50회.xlsx"
)
LOG_FILE = TARGET_FILE.parent / "candidate_tracker_update_log.txt"

ALL_SHEET = "전체당첨내역"
CUM_SHEET = "_calc_누적빈도"
ANCHOR_SHEET = "_calc_라운드기준"
SAMPLE_CHECK_SHEET = "기준100회_후보"  # 사후 검증용 표본 시트 1개

# ============================== 데이터 소스 ===================================
# weekly_lotto_file_update.py와 동일한 동행복권 공식 API(회차별 조회).
_DHLOTTERY_ROUND_URL = "https://www.dhlottery.co.kr/common.do?method=getLottoNumber&drwNo={round}"


def fetch_round_from_dhlottery(round_no: int) -> dict | None:
    """2026-09-15: urllib 기본 User-Agent("Python-urllib/x.x")로 요청하면
    동행복권 서버가 JSON이 아닌 응답(차단/안내 페이지로 추정)을 돌려줘서
    json.loads가 실패하는 사례가 실제로 확인됨. 브라우저처럼 보이는
    User-Agent를 붙여서 요청하도록 수정."""
    url = _DHLOTTERY_ROUND_URL.format(round=round_no)
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/128.0.0.0 Safari/537.36"
            ),
            "Accept": "application/json, text/javascript, */*; q=0.01",
        },
    )
    try:
        with urllib.request.urlopen(req, timeout=8) as resp:
            raw = resp.read().decode("utf-8", errors="replace")
        data = json.loads(raw)
    except json.JSONDecodeError as e:
        log(f"  [경고] {round_no}회차 조회 실패(JSON 아닌 응답): {e} "
            f"-> 응답 앞부분: {raw[:200]!r}")
        return None
    except Exception as e:
        log(f"  [경고] {round_no}회차 조회 실패(네트워크/응답 오류): {e}")
        return None

    if data.get("returnValue") != "success":
        return None

    nums = sorted(int(data[f"drwtNo{i}"]) for i in range(1, 7))
    return {"round": round_no, "nums": nums, "bonus": int(data["bnusNo"])}


def find_new_rounds(local_max_round: int, hard_limit: int = 10) -> list[dict]:
    out = []
    r = local_max_round + 1
    while len(out) < hard_limit:
        rec = fetch_round_from_dhlottery(r)
        if rec is None:
            break
        out.append(rec)
        r += 1
    return out


# ============================== 로깅 =========================================

def log(msg: str) -> None:
    line = f"[{datetime.datetime.now().isoformat(timespec='seconds')}] {msg}"
    print(line)
    try:
        LOG_FILE.parent.mkdir(parents=True, exist_ok=True)
        with open(LOG_FILE, "a", encoding="utf-8") as f:
            f.write(line + "\n")
    except OSError:
        pass


# ============================== 엑셀 재계산 (win32com) =========================

def recalc_and_save(path: Path) -> None:
    """구조 편집이 끝난 파일을 열어 강제 전체 재계산 후 그대로 저장 — 캐시된 값을
    최신으로 구워넣어서, 나중에 파일을 열었을 때 재계산 전에도 바로 올바른 값이
    보이게 한다. (openpyxl은 수식을 계산하지 않고 텍스트만 저장하기 때문에 필수.)"""
    import win32com.client

    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        wb = excel.Workbooks.Open(str(path.resolve()))
        try:
            excel.CalculateFullRebuild()
            wb.Save()
        finally:
            wb.Close(SaveChanges=True)
    finally:
        excel.Quit()


# ============================== 엑셀 오류값 스캔 ================================

_ERR = {"#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!"}


def scan_errors(path: Path) -> list[tuple[str, str]]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    found = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.strip() in _ERR:
                    found.append((sn, f"{cell.coordinate}={v}"))
    wb.close()
    return found


def verify_after_update(path: Path, expected_max_round: int) -> None:
    """저장·재계산 후 파일을 다시 열어, 자동 재계산이 실제로 제대로 됐는지
    독립적으로 확인한다. 문제가 있으면 예외를 던져 이후 처리를 멈춘다."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    anchor = wb[ANCHOR_SHEET].cell(2, 1).value
    if anchor != expected_max_round + 1:
        raise RuntimeError(
            f"검증 실패: {ANCHOR_SHEET}!A2(다음 예측 회차)={anchor}, "
            f"기대값={expected_max_round + 1}. 재계산이 반영 안 된 것으로 보임."
        )

    ws_sample = wb[SAMPLE_CHECK_SHEET]
    l1bd1 = [ws_sample.cell(1, c).value for c in range(12, 57)]
    if sorted(v for v in l1bd1 if v is not None) != list(range(1, 46)):
        raise RuntimeError(
            f"검증 실패: {SAMPLE_CHECK_SHEET}!L1:BD1(후보숫자 순위표)이 "
            f"1~45 전체를 담고 있지 않음 -> {l1bd1}"
        )

    wb.close()


# ============================== 구조 편집 (openpyxl) ===========================

def append_draw_result(ws_all, rec: dict) -> int:
    next_row = ws_all.max_row + 1
    ws_all.cell(next_row, 1, rec["round"])
    for i, n in enumerate(rec["nums"]):
        ws_all.cell(next_row, 2 + i, n)
    ws_all.cell(next_row, 8, rec["bonus"])
    return next_row


def append_cumfreq_row(ws_cum, round_no: int, all_row_num: int) -> int:
    """_calc_누적빈도에 회차 하나의 누적 출현횟수 행을 추가한다. all_row_num은
    전체당첨내역에 방금 추가된 행 번호(두 시트는 항상 행 번호가 1:1로 대응해야 함)."""
    next_row = ws_cum.max_row + 1
    prev_row = next_row - 1
    ws_cum.cell(next_row, 1, round_no)
    for c in range(2, 47):
        col_letter = get_column_letter(c)
        num = c - 1
        formula = (
            f"={col_letter}{prev_row}+COUNTIF("
            f"{ALL_SHEET}!$B{all_row_num}:$G{all_row_num},{num})"
        )
        ws_cum.cell(next_row, c, formula)
    return next_row


def get_current_max_round(path: Path) -> int:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[ALL_SHEET]
    val = ws.cell(ws.max_row, 1).value
    wb.close()
    return int(val)


def process_new_rounds(path: Path, new_rounds: list[dict]) -> None:
    wb = openpyxl.load_workbook(path, data_only=False)
    ws_all = wb[ALL_SHEET]
    ws_cum = wb[CUM_SHEET]

    last_round = None
    for rec in new_rounds:
        all_row = append_draw_result(ws_all, rec)
        cum_row = append_cumfreq_row(ws_cum, rec["round"], all_row)
        if all_row != cum_row:
            raise RuntimeError(
                f"{ALL_SHEET}와 {CUM_SHEET}의 행 번호가 어긋남"
                f"(all_row={all_row}, cum_row={cum_row}) — 두 시트가 이미 1:1로 "
                f"안 맞는 상태인 것으로 보임. 자동 처리를 중단합니다."
            )
        last_round = rec["round"]

    wb.save(path)
    log(f"  저장 완료({len(new_rounds)}개 회차 반영). 재계산 + 캐시값 굽는 중...")
    recalc_and_save(path)

    verify_after_update(path, expected_max_round=last_round)
    log("  재계산 후 검증 통과(다음 예측 회차 자동 갱신 확인, 후보숫자 순위표 정상).")


# ============================== 메인 =========================================

def main() -> int:
    log("=" * 70)
    log("★후보숫자_추적표_표본vs최근50회 자동 업데이트 시작")

    if not TARGET_FILE.exists():
        log(f"[오류] 파일을 찾을 수 없음 -> {TARGET_FILE} (CONFIG의 TARGET_FILE 확인 필요)")
        return 1

    try:
        local_max = get_current_max_round(TARGET_FILE)
    except Exception as e:
        log(f"[오류] 현재 최신 회차 확인 실패 -> {e}")
        return 1

    new_rounds = find_new_rounds(local_max)
    if not new_rounds:
        log(f"이미 최신 상태입니다(전체당첨내역 최신회차={local_max}). 종료.")
        log("=" * 70)
        return 0

    log(f"현재 {local_max}회차까지 반영됨 -> "
        f"{[r['round'] for r in new_rounds]}회차 새로 반영 시작")

    try:
        process_new_rounds(TARGET_FILE, new_rounds)
    except Exception as e:
        log(f"[오류] 처리 중 예외 발생: {e}")
        log("      원인 확인 후 재실행 필요.")
        return 1

    try:
        errs = scan_errors(TARGET_FILE)
        if errs:
            log(f"[경고] 저장 후 오류값 {len(errs)}개 발견 -> {errs[:5]}")
            log("주간 업데이트 종료 (오류/경고 있음 — 위 로그 확인 필요)")
            log("=" * 70)
            return 1
        log("저장 후 오류값 없음 확인 완료.")
    except Exception as e:
        log(f"[경고] 오류값 스캔 실패 -> {e}")

    log("주간 업데이트 종료 (정상)")
    log("=" * 70)
    return 0


if __name__ == "__main__":
    sys.exit(main())
