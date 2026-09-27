# -*- coding: utf-8 -*-
"""★조합생성_후보숫자_추적표 폴더 전체 파일 구조 진단(읽기 전용).

각 xlsx에 대해: 시트 목록 / 전체당첨내역 최신·최초 회차 / 각 '3차필터*' 시트의 K2·K3,
그리고 1행(L1:BD1)·2행·3행·4행 앞부분을 UTF-8 파일로 덤프한다.
"""
from __future__ import annotations

import glob
import os
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
TRACKER = ROOT / "★조합생성_후보숫자_추적표"
OUT = ROOT / "scratch" / "probe_tracker_files_out.txt"

lines: list[str] = []


def w(s: str = "") -> None:
    lines.append(s)


for path in sorted(glob.glob(str(TRACKER / "*.xlsx"))):
    name = os.path.basename(path)
    w("=" * 78)
    w(f"FILE: {name}")
    w(f"  mtime={__import__('datetime').datetime.fromtimestamp(os.path.getmtime(path)):%Y-%m-%d %H:%M}")
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:  # noqa: BLE001
        w(f"  !! load failed: {e}")
        continue
    w(f"  SHEETS({len(wb.sheetnames)}): {wb.sheetnames}")
    if "전체당첨내역" in wb.sheetnames:
        ws = wb["전체당첨내역"]
        rounds = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)
                  if isinstance(ws.cell(r, 1).value, int)]
        w(f"  전체당첨내역: rows={ws.max_row} cols={ws.max_column} "
          f"rounds={min(rounds) if rounds else None}~{max(rounds) if rounds else None} n={len(rounds)}")
    for sn in wb.sheetnames:
        if not (sn.startswith("3차필터") or sn.endswith("_후보")):
            continue
        ws = wb[sn]
        w(f"  ── sheet '{sn}' rows={ws.max_row} cols={ws.max_column}")
        w(f"     K1={ws.cell(1, 11).value!r} K2={ws.cell(2, 11).value!r} "
          f"K3={ws.cell(3, 11).value!r} K4={ws.cell(4, 11).value!r} J2={ws.cell(2, 10).value!r}")
        w(f"     A1..A12={[ws.cell(r, 1).value for r in range(1, 13)]}")
        for r in (1, 2, 3, 4):
            row = [ws.cell(r, c).value for c in range(12, 57)]
            w(f"     row{r}[L..BD][:8]={row[:8]} is_perm="
              f"{sorted(v for v in row if v is not None) == list(range(1, 46))}")
        # 순번 행(과거블록 직전 순열 행) 탐색
        block = next((r for r in range(2, ws.max_row + 1)
                      if isinstance(ws.cell(r, 1).value, int) and ws.cell(r, 1).value >= 1000), None)
        w(f"     first past-block row={block} A={ws.cell(block, 1).value if block else None} "
          f"B..G={[ws.cell(block, c).value for c in range(2, 9)] if block else None}")
        w(f"     L6 formula sample: {[ws.cell(6, c).value for c in (12, 13, 14, 15)]}")
        w(f"     L4 formula sample: {[ws.cell(4, c).value for c in (12, 13)]}")
        w(f"     L{block} row[:6]={[ws.cell(block, c).value for c in range(12, 18)] if block else None}")
        if block and block + 3 <= ws.max_row:
            w(f"     L{block+1} row[:6]={[ws.cell(block + 1, c).value for c in range(12, 18)]}")
    wb.close()

OUT.write_text("\n".join(lines), encoding="utf-8")
print(f"written {OUT} ({len(lines)} lines)")
