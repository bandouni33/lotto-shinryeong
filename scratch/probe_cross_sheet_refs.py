# -*- coding: utf-8 -*-
"""회차 컬럼 시트를 '참조하는 다른 셀'이 있는지 전수 조사 (읽기 전용).

왜 필요한가: 회차 컬럼을 밀면 가장 오래된 컬럼이 사라진다. 다른 시트의 수식이 그
시트/그 컬럼을 참조하고 있으면(openpyxl은 시트 간 참조를 자동 번역하지 않는다) 밀기
한 번에 파일이 조용히 어긋난다 — 자동화 전에 '참조 0건'을 실측으로 확정한다.

실행: venv312\\Scripts\\python.exe scratch\\probe_cross_sheet_refs.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as cl

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "★조합생성_후보숫자_추적표" / "조합생성_후보숫자_추적표_샘플.xlsx"
OUT = ROOT / "scratch" / "probe_cross_sheet_refs_out.txt"
TARGET_SHEETS = ("1차필터(7기본필터)", "2차필터(5이격수)")
L: list[str] = []


def w(s: str = "") -> None:
    L.append(str(s))
    print(s, flush=True)


wb = openpyxl.load_workbook(SRC, data_only=False)
try:
    total_formulas = 0
    hits: list[tuple[str, str, str]] = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                text = v.text if hasattr(v, "text") else v
                if not isinstance(text, str) or not text.startswith("="):
                    continue
                total_formulas += 1
                for tgt in TARGET_SHEETS:
                    if f"{tgt}!" in text or f"'{tgt}'!" in text:
                        hits.append((sn, c.coordinate, text[:120]))
                        break
    w(f"수식 셀 전체 {total_formulas:,}개")
    w(f"회차 컬럼 시트({TARGET_SHEETS})를 참조하는 셀: {len(hits)}개")
    for sn, coord, txt in hits[:20]:
        w(f"  {sn}!{coord} = {txt}")

    # 컬럼 시트 자체가 '다른 시트'를 참조하는 목록(밀어도 깨지지 않는지 확인용)
    for sn in TARGET_SHEETS:
        ws = wb[sn]
        refs: dict[str, int] = {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                text = v.text if hasattr(v, "text") else v
                if not isinstance(text, str) or not text.startswith("="):
                    continue
                for m in re.finditer(r"'?([^'!()]+)'?!\$?[A-Z]{1,3}\$?\d+", text):
                    name = m.group(1).strip()
                    if name and name != sn and not name.startswith(("IF", "OR", "AND", "SUM")):
                        refs[name] = refs.get(name, 0) + 1
        w(f"\n{sn} → 다른 시트 참조(괄호 제외): {refs}")

    # 회차 컬럼 범위 밖(예: A~M열)에서 회차 라벨을 참조하는 수식이 있는지
    for sn in TARGET_SHEETS:
        ws = wb[sn]
        first = min(c for c in range(1, ws.max_column + 1) if isinstance(ws.cell(4, c).value, int))
        outside = []
        for row in ws.iter_rows(min_col=1, max_col=first - 1):
            for c in row:
                v = c.value
                text = v.text if hasattr(v, "text") else v
                if isinstance(text, str) and text.startswith("=") and "$4" in text:
                    outside.append(f"{c.coordinate}={text[:80]}")
        w(f"{sn}: 회차 컬럼 왼쪽(A~{cl(first - 1)})에 '4행 라벨' 참조 수식 {len(outside)}개 {outside[:5]}")
finally:
    wb.close()

OUT.write_text("\n".join(L), encoding="utf-8")
print(f"written {OUT}")
sys.exit(1 if any(True for _ in []) else 0)
