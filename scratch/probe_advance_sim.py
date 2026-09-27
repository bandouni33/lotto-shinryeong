# -*- coding: utf-8 -*-
"""임시 복사본에서 '다음 회차 반영(weekly.process_one_round_for_file)'을 실제로 돌려 보고
소요 시간과 결과 구조를 확인한다(실물 파일은 건드리지 않는다).

엑셀 COM(재계산·캐시 굽기)은 테스트를 이 PC의 엑셀 상태에 매달리게 하므로 스텁한다
  · recalc_and_read_pending_row → 파이썬 폴백(_python_pending_forecast)
    (이 폴백 자체는 scratch/test_python_forecast_fallback.py가 엑셀 값과 일치함을 검증함)
  · recalc_and_save → no-op (openpyxl은 수식을 계산하지 않으므로 값은 엑셀이 담당)
"""
from __future__ import annotations

import hashlib
import os
import shutil
import sys
import tempfile
import time
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

import candidate_tracker_auto_update as tracker  # noqa: E402
import weekly_lotto_file_update as weekly  # noqa: E402

TRACKER = ROOT / "★조합생성_후보숫자_추적표"
SRC = TRACKER / "조합생성_후보숫자_추적표_전체표본_윈도우비교.xlsx"
OUT = ROOT / "scratch" / "probe_advance_sim_out.txt"
LINES: list[str] = []


def w(s: str = "") -> None:
    LINES.append(s)


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()[:16]


TMP = Path(tempfile.mkdtemp(prefix="lotto_advance_probe_"))
weekly.LOG_FILE = TMP / "probe_weekly.log"
tracker.LOG_FILE = TMP / "probe_tracker.log"
weekly.recalc_and_save = lambda path: None
weekly.recalc_and_read_pending_row = lambda path, sheets: weekly._python_pending_forecast(path, sheets)

before_src = sha(SRC)
copy = TMP / "copy.xlsx"
shutil.copy2(SRC, copy)
w(f"원본 sha(앞16)={before_src} → 복사본 {copy}")

t0 = time.time()
wb = openpyxl.load_workbook(copy, data_only=True, read_only=True)
draws = []
for row in wb["전체당첨내역"].iter_rows(min_row=2, max_col=8, values_only=True):
    if row[0] is None:
        continue
    draws.append((int(row[0]), [int(x) for x in row[1:7]]))
wb.close()
w(f"[{time.time() - t0:.1f}s] 이력 로드: {len(draws)}회차, 최신 {draws[-1][0]}")

latest = draws[-1][0]
rec = {"round": latest + 1, "nums": [2, 9, 17, 28, 36, 44], "bonus": 13}

t1 = time.time()
weekly.process_one_round_for_file(copy, "전체표본", rec)
w(f"[{time.time() - t1:.1f}s] process_one_round_for_file 완료(복사본)")

t2 = time.time()
wb2 = openpyxl.load_workbook(copy, data_only=True, read_only=True)
w(f"[{time.time() - t2:.1f}s] 결과 재로드")
try:
    ws_all = wb2["전체당첨내역"]
    rows = list(ws_all.iter_rows(min_row=2, max_col=8, values_only=True))
    w(f"전체당첨내역 마지막 두 행: {rows[-2:]}")
    for sn in [s for s in wb2.sheetnames if s.startswith("3차필터")]:
        ws = wb2[sn]
        n = int(sn.split("(")[1].split("회")[0])
        r7 = next(iter(ws.iter_rows(min_row=7, max_row=7, max_col=56, values_only=True)))
        r8 = next(iter(ws.iter_rows(min_row=8, max_row=8, max_col=56, values_only=True)))
        w2 = weekly._window_from_label(next(iter(ws.iter_rows(min_row=2, max_row=2, min_col=11, max_col=11, values_only=True)))[0])
        w3 = weekly._window_from_label(next(iter(ws.iter_rows(min_row=3, max_row=3, min_col=11, max_col=11, values_only=True)))[0])
        last = next(iter(ws.iter_rows(min_row=ws.max_row, max_row=ws.max_row, max_col=1, values_only=True)))[0]
        exp = tracker._prediction_for(draws, rec["round"], w2, w3)
        frozen = list(r8[11:56])
        w(f"  {sn}: max_row={ws.max_row} (기대 {6 + n}) A7={r7[0]} A8={r8[0]} 마지막행 A={last}")
        w(f"    8행 동결 예측 == 그 회차 직전 이력 재계산 ? {frozen == exp}")
        w(f"    8행 당첨번호 {list(r8[1:7])} / 적중수 I:K={list(r8[8:11])} (합 {sum(r8[8:11])})")
        w(f"    B7..K7 비어 있음 ? {all(v is None for v in r7[1:11])}")
finally:
    wb2.close()

TMP_HASH = sha(copy)
w(f"결과 복사본 sha={TMP_HASH}")
w(f"원본 sha 변화 없음 ? {sha(SRC) == before_src}")
w(f"총 경과 {time.time() - t0:.1f}s")

OUT.write_text("\n".join(LINES), encoding="utf-8")
print(f"written {OUT}")
