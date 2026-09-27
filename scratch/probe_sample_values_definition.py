# -*- coding: utf-8 -*-
"""샘플 파일의 '값' 칸들이 기존 검증된 정의로 재현되는지 실측(읽기 전용).

대상: 3차필터(100출현빈도순 후보) 시트의 K/L/M(상위·중위·하위 적중수).
가설: 그 값은 배포 엔진과 같은 정의 — 회차 R의 예측 = (R-1까지의 이력)으로 만든
      역대출현빈도 순위 − 최근100회 순위 격차순위(combo_filter_v2._gap_order_for_anchor)
      의 1~15위/16~30위/31~45위 그룹에 R의 당첨번호가 몇 개 들어갔는지.
검증: 과거행 각각에서 저장된 K/L/M 과 재계산값을 대조(불일치 0이면 정의 확정).

실행: venv312\\Scripts\\python.exe scratch\\probe_sample_values_definition.py
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

import combo_filter_v2 as cf  # noqa: E402

SRC = ROOT / "★조합생성_후보숫자_추적표" / "조합생성_후보숫자_추적표_샘플.xlsx"
OUT = ROOT / "scratch" / "probe_sample_values_definition_out.txt"
LINES: list[str] = []


def w(s: str = "") -> None:
    LINES.append(s)


wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
try:
    ws_all = wb["전체당첨내역"]
    draws = []
    for row in ws_all.iter_rows(min_row=2, max_col=8, values_only=True):
        if row[0] is None:
            continue
        draws.append((int(row[0]), [int(x) for x in row[1:7]], int(row[7] or 0)))
    draws.sort(key=lambda t: t[0])
    hist = [{"draw_round": r, "nums": ns, "bonus": b} for r, ns, b in draws]
    by_round = {r: (ns, b) for r, ns, b in draws}
    w(f"이력 {len(draws)}회차, 최신 {draws[-1][0]}")

    for sn in wb.sheetnames:
        if not any(k in sn for k in ("3차필터", "회차별 추적표", "추적결과", "4차필터")):
            continue
        ws = wb[sn]
        w(f"\n=== [{sn}] rows={ws.max_row} cols={ws.max_column}")
        matched = mismatch = skipped = 0
        for r in range(3, ws.max_row + 1):
            rnd = ws.cell(r, 1).value
            nums = [ws.cell(r, c).value for c in range(2, 8)]
            got = [ws.cell(r, c).value for c in (11, 12, 13)]      # K/L/M
            if not isinstance(rnd, int) or any(v is None for v in nums) or any(v is None for v in got):
                skipped += 1
                continue
            if rnd - 1 not in by_round:
                skipped += 1
                continue
            order = cf._gap_order_for_anchor(hist, rnd - 1)
            exp = [sum(1 for n in nums if n in set(order[i * 15:(i + 1) * 15])) for i in range(3)]
            ok = list(got) == exp
            matched += ok
            mismatch += (not ok)
            if not ok and mismatch <= 6:
                w(f"  {r}행 {rnd}회차: 시트 K/L/M={got} vs 재계산 {exp}  당첨 {nums}")
        w(f"  검사 {matched + mismatch}행(건너뜀 {skipped}) · 일치 {matched} · 불일치 {mismatch}")
finally:
    wb.close()

OUT.write_text("\n".join(LINES), encoding="utf-8")
print(f"written {OUT}")
