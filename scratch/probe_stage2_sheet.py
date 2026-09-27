# -*- coding: utf-8 -*-
"""2차필터(5이격수) / 2차추적결과 / 4차필터 / 3차필터 구조 정밀 덤프 (읽기 전용).

사용자 확정(2026-09-27): 2차필터 J열 = '이격수'(로또 번호가 아님).
여기서는 그 시트의 회차열 수식 원문과 헬퍼행(1500대)을 그대로 떠서
"이격수를 어떻게 세는지"를 파일 자체 근거로 확정한다.

실행: venv312\\Scripts\\python.exe scratch\\probe_stage2_sheet.py
"""
from __future__ import annotations

import os
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)

SRC = ROOT / "★조합생성_후보숫자_추적표" / "조합생성_후보숫자_추적표_샘플.xlsx"
OUT = ROOT / "scratch" / "probe_stage2_sheet_out.txt"
L: list[str] = []


def w(s: str = "") -> None:
    L.append(str(s))


def txt(v, n: int = 500) -> str:
    if v is None:
        return "None"
    if hasattr(v, "text"):
        t = v.text
        return t if len(t) <= n else t[:n] + f"…(+{len(t) - n}자)"
    s = str(v)
    return s if len(s) <= n else s[:n] + "…"


wv = openpyxl.load_workbook(SRC, data_only=True)
wf = openpyxl.load_workbook(SRC, data_only=False)
try:
    w(f"시트 목록: {wv.sheetnames}")
    for sn in wv.sheetnames:
        w(f"  {sn}: rows={wv[sn].max_row} cols={wv[sn].max_column}")

    for name in ("2차필터(5이격수)", "1차필터(7기본필터)"):
        ws, wsf = wv[name], wf[name]
        w(f"\n=== {name}: 1~4행 전 셀 ===")
        for r in range(1, 5):
            cells = [f"{openpyxl.utils.get_column_letter(c)}{r}={txt(ws.cell(r, c).value, 60)}"
                     for c in range(1, min(ws.max_column, 20) + 1)
                     if ws.cell(r, c).value is not None]
            w(f"  row{r}: " + " | ".join(cells))
        w(f"\n=== {name}: 규칙행 H/J/K/L + 첫 회차열 수식 (5~40행) ===")
        first_round_col = None
        for c in range(5, ws.max_column + 1):
            if isinstance(ws.cell(4, c).value, int):
                first_round_col = c
                break
        colL = openpyxl.utils.get_column_letter(first_round_col) if first_round_col else "?"
        w(f"  첫 회차열 = {colL} (4행값 {ws.cell(4, first_round_col).value if first_round_col else None})")
        for r in range(5, 41):
            h, j, k, lm = (ws.cell(r, 8).value, ws.cell(r, 10).value,
                           ws.cell(r, 11).value, ws.cell(r, 12).value)
            if h is None and j is None and k is None and lm is None:
                continue
            f = wsf.cell(r, first_round_col).value if first_round_col else None
            w(f"  row{r}: H={txt(h, 30)} J={txt(j, 80)} K={k} L={lm}")
            if f is not None:
                w(f"        {colL}{r} = {txt(f, 700)}")

        w(f"\n=== {name}: 헬퍼행 1500~1520, A~{openpyxl.utils.get_column_letter(min(ws.max_column, 14))}열 (값) ===")
        for r in range(1500, 1521):
            cells = []
            for c in range(1, min(ws.max_column, 14) + 1):
                v = ws.cell(r, c).value
                if v is not None:
                    cells.append(f"{openpyxl.utils.get_column_letter(c)}={txt(v, 40)}")
            if cells:
                w(f"  row{r}: " + " | ".join(cells))
        w(f"\n=== {name}: 헬퍼행 1500~1520 첫 회차열 수식 ===")
        if first_round_col:
            for r in range(1500, 1521):
                f = wsf.cell(r, first_round_col).value
                if f is not None:
                    w(f"  {colL}{r} = {txt(f, 700)}")

    for name in ("2차추적결과", "1차추적결과", "4차필터", "3차필터(100출현빈도순 후보)"):
        ws, wsf = wv[name], wf[name]
        w(f"\n=== {name}: rows={ws.max_row} cols={ws.max_column} ===")
        for r in range(1, 4):
            cells = [f"{openpyxl.utils.get_column_letter(c)}{r}={txt(ws.cell(r, c).value, 45)}"
                     for c in range(1, min(ws.max_column, 26) + 1)
                     if ws.cell(r, c).value is not None]
            w(f"  값 row{r}: " + " | ".join(cells))
        for r in range(1, 4):
            cells = [f"{openpyxl.utils.get_column_letter(c)}{r}={txt(wsf.cell(r, c).value, 300)}"
                     for c in range(1, min(ws.max_column, 26) + 1)
                     if wsf.cell(r, c).value is not None]
            w(f"  식 row{r}: " + " | ".join(cells))
        w(f"  3행 이후 회차 라벨(A열) 5개: "
          f"{[ws.cell(r, 1).value for r in range(3, 8)]}")
finally:
    wv.close()
    wf.close()

OUT.write_text("\n".join(L), encoding="utf-8")
print(f"written {OUT} ({len(L)}줄)")
