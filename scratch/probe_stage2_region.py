# -*- coding: utf-8 -*-
"""2차필터(5이격수) 시트의 41행~끝 영역, 3차필터 후보영역, 회차별 추적표 구조 덤프(읽기 전용).

배경: 2차 시트 회차열에는 규칙 수식이 없고 리터럴(1236 등)만 있었다.
규칙행이 어디까지인지, 계산 영역이 어디인지 파일 자체 근거로 확정한다.

실행: venv312\\Scripts\\python.exe scratch\\probe_stage2_region.py
"""
from __future__ import annotations

import os
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as cl

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)

SRC = ROOT / "★조합생성_후보숫자_추적표" / "조합생성_후보숫자_추적표_샘플.xlsx"
OUT = ROOT / "scratch" / "probe_stage2_region_out.txt"
L: list[str] = []


def w(s: str = "") -> None:
    L.append(str(s))


def txt(v, n: int = 90) -> str:
    if v is None:
        return "None"
    if hasattr(v, "text"):
        t = v.text
        return t if len(t) <= n else t[:n] + "…"
    s = str(v)
    return s if len(s) <= n else s[:n] + "…"


wv = openpyxl.load_workbook(SRC, data_only=True)
wf = openpyxl.load_workbook(SRC, data_only=False)
try:
    name = "2차필터(5이격수)"
    ws, wsf = wv[name], wf[name]
    w(f"=== {name}: rows={ws.max_row} cols={ws.max_column} ===")
    w("\n-- 행별 요약(값 있는 셀 수 / 수식 셀 수, 41행~) --")
    for r in range(41, ws.max_row + 1):
        vals = [(c, ws.cell(r, c).value) for c in range(1, ws.max_column + 1)
                if ws.cell(r, c).value is not None]
        fs = [c for c in range(1, ws.max_column + 1) if wsf.cell(r, c).value is not None
              and str(wsf.cell(r, c).value).startswith("=")]
        if vals or fs:
            w(f"  row{r}: 값{len(vals)}개 수식{len(fs)}개 A열={txt(ws.cell(r,1).value,20)} "
              f"H={txt(ws.cell(r,8).value,25)} J={txt(ws.cell(r,10).value,40)} "
              f"K={ws.cell(r,11).value} L={ws.cell(r,12).value}")

    w("\n-- 값 있는 셀 전량 (41행~, 열문자=값) --")
    shown = 0
    for r in range(41, ws.max_row + 1):
        cells = [f"{cl(c)}{r}={txt(ws.cell(r, c).value, 60)}"
                 for c in range(1, ws.max_column + 1)
                 if ws.cell(r, c).value is not None]
        if cells:
            w("  " + " | ".join(cells))
            shown += 1
        if shown > 60:
            w(f"  …(이하 생략, 총 표시 {shown}행)")
            break

    w("\n-- 규칙행 개수 집계(2차) --")
    n_like = 0
    for r in range(5, ws.max_row + 1):
        j, k, lm = ws.cell(r, 10).value, ws.cell(r, 11).value, ws.cell(r, 12).value
        if j is not None and isinstance(k, (int, float)) and isinstance(lm, (int, float)):
            n_like += 1
    w(f"  J+K+L 모두 있는 행 = {n_like}")
    w("  5~60행 J/K/L:")
    for r in range(5, 61):
        j, k, lm = ws.cell(r, 10).value, ws.cell(r, 11).value, ws.cell(r, 12).value
        if j is not None or k is not None or lm is not None:
            w(f"    row{r}: H={txt(ws.cell(r,8).value,22)} J={txt(j,70)} K={k} L={lm}")

    for name in ("3차필터(100출현빈도순 후보)", "회차별 추적표"):
        ws, wsf = wv[name], wf[name]
        w(f"\n=== {name}: rows={ws.max_row} cols={ws.max_column} ===")
        for r in range(1, min(12, ws.max_row) + 1):
            cells = [f"{cl(c)}{r}={txt(ws.cell(r, c).value, 28)}"
                     for c in range(1, min(ws.max_column, 26) + 1)
                     if ws.cell(r, c).value is not None]
            w(f"  값 row{r}: " + " | ".join(cells))
        w("  -- 수식 있는 셀(앞 25개) --")
        cnt = 0
        for r in range(1, min(ws.max_row, 150) + 1):
            for c in range(1, ws.max_column + 1):
                f = wsf.cell(r, c).value
                if f is not None and str(f).startswith("="):
                    w(f"    {cl(c)}{r} = {txt(f, 200)}")
                    cnt += 1
                if cnt >= 25:
                    break
            if cnt >= 25:
                break
finally:
    wv.close()
    wf.close()

OUT.write_text("\n".join(L), encoding="utf-8")
print(f"written {OUT} ({len(L)}줄)")
