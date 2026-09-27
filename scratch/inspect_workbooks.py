# -*- coding: utf-8 -*-
"""★조합생성_후보숫자_추적표 워크북 5개 구조·수식 조사기(읽기 전용).

사용:
    python scratch\\inspect_workbooks.py                 # 시트 구조 + 키워드 셀 위치
    python scratch\\inspect_workbooks.py 셀  샘플 1차필터 A1:N20   # 특정 시트 범위의 수식 원문
    python scratch\\inspect_workbooks.py 값  샘플 1차필터 A1:N20   # 같은 범위의 캐시값
"""

from __future__ import annotations

import os
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
TRACKER = ROOT / "★조합생성_후보숫자_추적표"

FILES = {
    "샘플": TRACKER / "조합생성_후보숫자_추적표_샘플.xlsx",
    "200회검증용": TRACKER / "조합생성_후보숫자_추적표_샘플_200회검증용.xlsx",
    "전체표본": TRACKER / "조합생성_후보숫자_추적표_전체표본_윈도우비교.xlsx",
    "최근500표본": TRACKER / "조합생성_후보숫자_추적표_최근500표본_윈도우비교.xlsx",
    "표본vs최근50회": TRACKER / "★후보숫자_추적표_표본vs최근50회.xlsx",
}

KEYWORDS = ("잔량", "순번", "AUTO", "1차필터", "2차필터", "3차필터", "4차", "필터")


def sheet_map() -> None:
    for label, path in FILES.items():
        wb = openpyxl.load_workbook(path, data_only=False, read_only=True)
        print(f"\n=== {label} ({path.name}) ===")
        for sn in wb.sheetnames:
            ws = wb[sn]
            state = "숨김" if ws.sheet_state != "visible" else ""
            print(f"   {sn!r:38} {ws.max_row:6}행 x {ws.max_column:3}열 {state}")
        wb.close()


def keyword_cells(limit_per_sheet: int = 12) -> None:
    for label, path in FILES.items():
        wb = openpyxl.load_workbook(path, data_only=False)
        print(f"\n=== {label} ===")
        for sn in wb.sheetnames:
            ws = wb[sn]
            hits = []
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    text = v.text if hasattr(v, "text") else v  # ArrayFormula
                    if isinstance(text, str) and any(k in text for k in KEYWORDS):
                        hits.append((c.coordinate, text.replace("\n", " ")[:110]))
            if hits:
                print(f"  [{sn}] 키워드 셀 {len(hits)}개")
                for coord, text in hits[:limit_per_sheet]:
                    print(f"     {coord:>8} {text}")
                if len(hits) > limit_per_sheet:
                    print(f"     ... 외 {len(hits) - limit_per_sheet}개")
        wb.close()


def dump(label: str, sheet: str, rng: str, formulas: bool) -> None:
    wb = openpyxl.load_workbook(FILES[label], data_only=not formulas)
    ws = wb[sheet]
    print(f"=== {label} / {sheet} / {rng} / {'수식' if formulas else '캐시값'} ===")
    for row in ws[rng]:
        cells = []
        for c in row:
            v = c.value
            if hasattr(v, "text"):
                v = "{" + str(v.text)[:90] + "}"
            elif isinstance(v, str) and len(v) > 90:
                v = v[:90] + "…"
            cells.append(f"{c.coordinate}={v!r}")
        line = "  ".join(cells)
        if line.strip().replace("=None", "").strip():
            print("  " + line)
    wb.close()


def all_strings(label: str, cap: int = 70) -> None:
    """워크북 전체에서 '문자열인 셀'을 (좌표, 값)으로 모아 보여준다 — 헤더/라벨
    어휘를 파악해서 '필터잔량' 같은 이름이 실제로 어느 셀에 있는지 찾기 위한 모드."""
    wb = openpyxl.load_workbook(FILES[label], data_only=False)
    for sn in wb.sheetnames:
        ws = wb[sn]
        seen: dict[str, str] = {}
        for row in ws.iter_rows():
            for c in row:
                text = c.value.text if hasattr(c.value, "text") else c.value
                if isinstance(text, str) and text.strip():
                    key = text.replace("\n", " ")[:70]
                    seen.setdefault(key, c.coordinate)
        if not seen:
            continue
        print(f"\n=== {label} / {sn} — 문자열 셀 {len(seen)}종 ===")
        for text, coord in list(seen.items())[:cap]:
            print(f"   {coord:>8} {text!r}")
        if len(seen) > cap:
            print(f"   ... 외 {len(seen) - cap}종")
    wb.close()


if __name__ == "__main__":
    args = sys.argv[1:]
    if not args:
        sheet_map()
        keyword_cells()
    elif args[0] == "셀" and len(args) == 4:
        dump(args[1], args[2], args[3], formulas=True)
    elif args[0] == "값" and len(args) == 4:
        dump(args[1], args[2], args[3], formulas=False)
    elif args[0] == "문자" and len(args) == 2:
        all_strings(args[1])
    else:
        print(__doc__)
