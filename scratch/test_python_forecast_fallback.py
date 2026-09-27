# -*- coding: utf-8 -*-
"""`weekly_lotto_file_update._python_pending_forecast` (엑셀 없는 경로) 검증.

2026-09-27 작업에서 추가한 폴백인데, 실제 실행에서는 **한 번도 타지 않았다**
(3차필터 파일들은 엑셀이 성공했고, 엑셀이 실패한 파일은 3차필터 시트가 없는
샘플·200회검증용뿐이었다). 실행되지 않은 코드는 검증되지 않은 코드이므로,
"엑셀이 계산해 둔 값"과 직접 맞대어 본다 — 그게 이 폴백의 유일한 계약이다.

불변식(모든 유효 입력에 대해):
  A) 시트 라벨 → 창(window) 파싱: '전체'/'(최근500회)'/'(최근N회)'/'최근50회 빈도(고정)'/
     빈 문자열·None 전부에서 정확할 것
  B) 어떤 시트에 대해서도 결과는 1~45의 순열일 것
  C) 결과는 워크북 3차필터 시트의 7행(엑셀이 계산해 둔 다음회차 순번 1~45)과
     원소 순서까지 완전히 같을 것 — 즉 엑셀 없이도 같은 값을 낸다

전부 읽기 전용(워크북·DB 수정 없음):
    venv312\\Scripts\\python.exe scratch\\test_python_forecast_fallback.py
"""

from __future__ import annotations

import os
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

import weekly_lotto_file_update as weekly  # noqa: E402

TRACKER = ROOT / "★조합생성_후보숫자_추적표"
NUMS = list(range(1, 46))


def setUpModule():
    """테스트가 실제 로그(weekly_update_log.txt)를 오염시키지 않게 임시 경로로 돌린다."""
    weekly.LOG_FILE = Path(tempfile.gettempdir()) / "lotto_fallback_test.log"


class WindowFromLabelTests(unittest.TestCase):
    """A) 라벨 표기 변형 전체 — 이 파싱이 틀리면 폴백 전체가 조용히 다른 창을 쓴다."""

    def test_known_labels(self):
        cases = [
            ("기준빈도(전체)", None),                 # 전체표본 2행
            ("기준빈도(최근500회)", 500),             # 최근500표본 2행
            ("최근100회 빈도", 100),                  # 전체표본 3행
            ("최근500회 빈도", 500),
            ("기준빈도(최근100회)", 100),             # 표본vs최근50회 2행
            ("기준빈도 회차수(이 값만 바꾸면 전체 재계산)", None),  # A1 안내문(창 아님)
            ("", None),
            (None, None),
            ("최근 200 회 빈도", 200),                # 공백 변형
        ]
        for label, expected in cases:
            with self.subTest(label=label):
                self.assertEqual(weekly._window_from_label(label), expected)

    def test_default_is_returned_when_no_window_in_label(self):
        self.assertEqual(weekly._window_from_label("기준빈도(전체)", default=7), 7)
        self.assertEqual(weekly._window_from_label(None, default=7), 7)
        self.assertEqual(weekly._window_from_label("숫자없음", default=7), 7)


class FallbackMatchesExcelTests(unittest.TestCase):
    """B·C) 워크북 3차필터 12개 시트 전부에서 엑셀 재계산값과 완전히 일치해야 한다."""

    FILES = {
        "전체표본": TRACKER / "조합생성_후보숫자_추적표_전체표본_윈도우비교.xlsx",
        "최근500표본": TRACKER / "조합생성_후보숫자_추적표_최근500표본_윈도우비교.xlsx",
    }

    def _three_cha_sheets(self, path: Path, data_only: bool) -> tuple[list[str], object]:
        wb = openpyxl.load_workbook(path, data_only=data_only, read_only=True)
        return [sn for sn in wb.sheetnames if sn.startswith("3차필터")], wb

    def test_every_sheet_matches_excel_calculated_row(self):
        checked = 0
        for label, path in self.FILES.items():
            sheets, wb = self._three_cha_sheets(path, data_only=False)
            wb.close()
            self.assertTrue(sheets, f"{label}: 3차필터 시트가 하나도 없다 — 테스트 전제 붕괴")

            got = weekly._python_pending_forecast(path, sheets)

            _, wb_cached = self._three_cha_sheets(path, data_only=True)
            try:
                for sn in sheets:
                    cached = [wb_cached[sn].cell(7, c).value for c in range(12, 57)]
                    with self.subTest(file=label, sheet=sn):
                        # B) 순열 불변식
                        self.assertEqual(sorted(got[sn]), NUMS, f"{sn}: 결과가 1~45 순열이 아님")
                        # C) 엑셀과 완전 일치 (캐시가 없으면 대조 자체가 불가 — 실패로 알린다)
                        self.assertEqual(sorted(v for v in cached if v is not None), NUMS,
                                         f"{sn}: 엑셀 7행 캐시값이 없어 대조 불가(먼저 엑셀로 재계산 필요)")
                        self.assertEqual(list(got[sn]), list(cached),
                                         f"{sn}: 파이썬 폴백이 엑셀 값과 다름")
                    checked += 1
            finally:
                wb_cached.close()
        self.assertEqual(checked, 12, f"검사한 시트 수가 12가 아니다({checked})")

    def test_no_new_rounds_means_same_result(self):
        """앞선 호출이 워크북을 건드리지 않는다 — 두 번 불러도 같은 값."""
        for label, path in self.FILES.items():
            sheets, wb = self._three_cha_sheets(path, data_only=False)
            wb.close()
            first = weekly._python_pending_forecast(path, sheets)
            second = weekly._python_pending_forecast(path, sheets)
            self.assertEqual(first, second, f"{label}: 반복 호출 결과가 달라짐(파일 수정 또는 비결정성)")


if __name__ == "__main__":
    unittest.main(verbosity=2)
