# -*- coding: utf-8 -*-
"""샘플 파일 컬럼식 구조 정밀 실측(읽기 전용) — 컬럼 확장 자동화 설계용.

알아내려는 것:
  · 각 시트의 '회차 축'이 행인지 열인지, 라벨이 어디 있고 몇 개가 있는지
  · 1차필터(7기본필터)·2차필터(5이격수) 시트에서 '회차 열 하나'가 어떤 셀로 구성되는지
    (수식이 어떤 주소를 참조하는지, 스테이징 블록 1504~1516행이 무엇을 하는지)
  · 1차/2차 추적결과·3차필터·4차필터·회차별 추적표가 '회차 행'을 어떻게 기록하는지

출력은 비어있지 않은 셀만 r,c=값 형태로 압축한다.
실행: venv312\\Scripts\\python.exe scratch\\probe_sample_structure.py
"""
from __future__ import annotations

import os
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
SRC = ROOT / "★조합생성_후보숫자_추적표" / "조합생성_후보숫자_추적표_샘플.xlsx"
OUT = ROOT / "scratch" / "probe_sample_structure_out.txt"
LINES: list[str] = []


def w(s: str = "") -> None:
    LINES.append(s)


def dump_grid(ws, r0, r1, c0, c1, label, cap=60) -> None:
    """지정 구역의 비어있지 않은 셀을 r,c=값(수식은 앞 60자)으로 압축 출력."""
    n = 0
    for row in ws.iter_rows(min_row=r0, max_row=r1, min_col=c0, max_col=c1):
        for c in row:
            v = c.value
            t = v.text if hasattr(v, "text") else v
            if t is None:
                continue
            s = str(t)
            if len(s) > 60:
                s = s[:60] + "…"
            w(f"    {c.coordinate}={s!r}")
            n += 1
            if n >= cap:
                w("    …(생략)")
                return


# ── 값 보기 ──────────────────────────────────────────────────────────────
wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
try:
    w(f"SRC: {SRC.name} (값 보기)")
    for sn in wb.sheetnames:
        ws = wb[sn]
        w(f"\n=== [{sn}] rows={ws.max_row} cols={ws.max_column}")
        dump_grid(ws, 1, 6, 1, 14, "머리글 1~6행 × A~N")
        if sn.startswith(("1차필터", "2차필터")):
            # 회차 라벨 행(값 4행 부근) 전체 열
            w("    -- 회차 라벨 행(4행) 전체 열:")
            dump_grid(ws, 4, 5, 1, ws.max_column, "row4", cap=45)
            w("    -- 스테이징 구역 1495~1516행 × A~H:")
            dump_grid(ws, 1495, 1516, 1, 8, "staging", cap=60)
        if sn.endswith("추적결과") or sn == "4차필터":
            w("    -- 3행(최신 회차행) 전체 열:")
            dump_grid(ws, 3, 3, 1, ws.max_column, "row3", cap=25)
            w(f"    -- 마지막 행({ws.max_row}행) 전체 열:")
            dump_grid(ws, ws.max_row, ws.max_row, 1, ws.max_column, "last", cap=25)
        if "3차필터" in sn or "회차별 추적표" in sn:
            dump_grid(ws, 6, 9, 1, 24, "6~9행 × A~X", cap=50)
finally:
    wb.close()

# ── 수식 보기 ────────────────────────────────────────────────────────────
wf = openpyxl.load_workbook(SRC, data_only=False, read_only=True)
try:
    w(f"\n\nSRC: {SRC.name} (수식 보기)")
    for sn in wf.sheetnames:
        ws = wf[sn]
        fcount = 0
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if hasattr(v, "text") or (isinstance(v, str) and v.startswith("=")):
                    fcount += 1
        w(f"\n=== [{sn}] rows={ws.max_row} cols={ws.max_column} 수식셀={fcount}")
        if sn.startswith(("1차필터", "2차필터")):
            w("    -- 첫 데이터 열(N, 14열) 1~14행 수식:")
            dump_grid(ws, 1, 14, 14, 14, "colN", cap=30)
            w("    -- 둘째 데이터 열(O, 15열) 1~14행 수식:")
            dump_grid(ws, 1, 14, 15, 15, "colO", cap=30)
            w("    -- A~M 1~14행(행머리글) 수식/값:")
            dump_grid(ws, 1, 14, 1, 13, "ABM", cap=40)
            w("    -- 1495~1516행 × A~F 수식(스테이징):")
            dump_grid(ws, 1495, 1516, 1, 6, "staging", cap=40)
        if sn.endswith("추적결과") or sn == "4차필터":
            w("    -- 1~4행 × A~S(수식):")
            dump_grid(ws, 1, 4, 1, 19, "1-4", cap=40)
            w("    -- 마지막 두 행 × A~S(수식):")
            dump_grid(ws, ws.max_row - 1, ws.max_row, 1, 19, "last2", cap=20)
        if "3차필터" in sn:
            w("    -- 1~9행 × A~N(수식):")
            dump_grid(ws, 1, 9, 1, 14, "1-9", cap=50)
        if "회차별 추적표" in sn:
            w("    -- 1~6행 × A~N(수식):")
            dump_grid(ws, 1, 6, 1, 14, "1-6", cap=50)
finally:
    wf.close()

OUT.write_text("\n".join(LINES), encoding="utf-8")
print(f"written {OUT} ({len(LINES)} lines)")
