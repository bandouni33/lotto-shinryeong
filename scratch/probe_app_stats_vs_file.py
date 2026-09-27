# -*- coding: utf-8 -*-
"""파일 추적결과 열 vs 앱 자동화 기록(draw_generation_stats) 대조 (읽기 전용).

목적: 사용자 목표가 "매주 배포용 조합생성 수량이 앱 자동화 수치와 일치하는지 검증하는
대조표"이므로, 먼저 **파일의 어느 열이 앱 자동화 수치인지** 실측으로 확정한다.
   · 파일 1차추적결과 I(최종) / 2차추적결과 I / 4차필터 I  ← 후보
   · DB draw_generation_stats.stage2_count / stage4_count / rank1~5

DB는 SELECT만 한다(쓰기·마이그레이션 없음). 테이블 생성 플래그를 세워
marketing_db가 초기화 DDL을 실행하지 않게 막는다.

실행: venv312\\Scripts\\python.exe scratch\\probe_app_stats_vs_file.py
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

SRC = ROOT / "★조합생성_후보숫자_추적표" / "조합생성_후보숫자_추적표_샘플.xlsx"
OUT = ROOT / "scratch" / "probe_app_stats_vs_file_out.txt"
L: list[str] = []


def w(s: str = "") -> None:
    L.append(str(s))
    print(s, flush=True)


# ── 1) 파일 추적결과 열
wb = openpyxl.load_workbook(SRC, data_only=True)
try:
    file_rows: dict[int, dict] = {}
    for sh, key in (("1차추적결과", "s1"), ("2차추적결과", "s2"), ("4차필터", "s4")):
        ws = wb[sh]
        for r in range(3, ws.max_row + 1):
            rr = ws.cell(r, 1).value
            if not isinstance(rr, int):
                continue
            d = file_rows.setdefault(rr, {"nums": [ws.cell(r, c).value for c in range(2, 8)]})
            d[key] = ws.cell(r, 9).value
            d[key + "_H"] = ws.cell(r, 8).value
            if key == "s1":
                d["S"] = ws.cell(r, 19).value
finally:
    wb.close()
w(f"파일 추적결과 회차 = {len(file_rows)}개 ({min(file_rows)}~{max(file_rows)})")

# ── 2) DB 앱 자동화 기록
from env_loader import load_dotenv_file  # noqa: E402

load_dotenv_file()

import marketing_db  # noqa: E402

marketing_db._MARKETING_TABLES_READY = True      # 초기화 DDL 실행 방지
conn = marketing_db._connect()
try:
    rows = conn.execute(
        "SELECT draw_round, stage2_count, stage4_count, rank1_num, rank2_num, rank3_num, "
        "rank4_num, rank5_num, recorded_at FROM draw_generation_stats ORDER BY draw_round"
    ).fetchall()
finally:
    conn.close()

db: dict[int, dict] = {}
for r in rows:
    try:
        db[int(r[0])] = {"stage2": r[1], "stage4": r[2],
                         "top": [r[3], r[4], r[5], r[6], r[7]], "at": r[8]}
    except (TypeError, IndexError):
        db[int(r["draw_round"])] = {"stage2": r["stage2_count"], "stage4": r["stage4_count"],
                                    "top": [r["rank1_num"], r["rank2_num"], r["rank3_num"],
                                            r["rank4_num"], r["rank5_num"]],
                                    "at": r["recorded_at"]}
w(f"DB draw_generation_stats 행 = {len(db)}개"
  + (f" ({min(db)}~{max(db)})" if db else ""))

# ── 3) 대조표
w("\n== 대조 (파일 열 vs DB 기록) ==")
w(f"{'회차':>6} {'파일1차I':>11} {'파일2차I':>11} {'파일4차I':>11} "
  f"{'DB_stage2':>11} {'DB_stage4':>11} {'파일2차==db2':>12} {'파일4차==db4':>12}")
common = sorted(set(db) & set(file_rows), reverse=True)
for rnd in common:
    f, d = file_rows[rnd], db[rnd]
    m2 = (f.get("s2") == d["stage2"])
    m4 = (f.get("s4") == d["stage4"])
    w(f"{rnd:>6} {f.get('s1', 0):>11,} {f.get('s2', 0):>11,} {f.get('s4', 0):>11,} "
      f"{d['stage2']:>11,} {d['stage4']:>11,} {str(m2):>12} {str(m4):>12}")
w(f"\n공통 회차 {len(common)}개 · 파일에만 {len(set(file_rows) - set(db))}개 · DB에만 {len(set(db) - set(file_rows))}개")
w(f"DB에만 있는 회차: {sorted(set(db) - set(file_rows), reverse=True)[:20]}")

# ── 4) 격차순위 top5 대조 (파일 K~M 상중하 대신 top5는 3차필터 시트에 있다)
w("\n== DB 기록된 격차순위 top5 ==")
for rnd in common[:6]:
    w(f"  {rnd}회차: {db[rnd]['top']} (recorded_at {db[rnd]['at']})")

# ── 5) 일치 요약
m2_all = sum(1 for r in common if file_rows[r].get("s2") == db[r]["stage2"])
m4_all = sum(1 for r in common if file_rows[r].get("s4") == db[r]["stage4"])
w(f"\n2차추적결과 I == DB stage2_count : {m2_all}/{len(common)}")
w(f"4차필터 I == DB stage4_count     : {m4_all}/{len(common)}")

Path(OUT).write_text("\n".join(L), encoding="utf-8")
print(f"written {OUT}")
os._exit(0)
