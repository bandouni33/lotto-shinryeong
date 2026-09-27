# -*- coding: utf-8 -*-
"""3차필터 블록 정합화(발견 B 대응)의 불변식 검증.

검증 대상: candidate_tracker_auto_update.py의
  · _prediction_for()          — 그 회차 직전 이력만으로 격차순위를 만든다
  · realign_sliding_blocks()   — 밀린 블록의 예측순위(L:BD)·적중수(I:K)를 다시 쓴다
  · _block_mismatch_rows()     — 어긋난 행을 (읽기 전용으로) 세어낸다

불변식(모든 유효 입력에 대해):
  A) 예측은 미래를 보지 않는다: 회차 X의 예측은 X 이후 데이터가 있든 없든 같다
  B) 예측은 항상 1~45 순열이다
  C) 정합화 후 각 행은 자기 회차의 예측과 자기 회차의 당첨번호를 짝지으므로,
     상위/중위/하위 적중수가 재계산과 정확히 같고 합은 6이다
  D) 정합화는 멱등적이다: 두 번 적용해도 결과가 같고, 그 뒤 감지 결과는 0이다
  E) 어긋난 입력을 넣으면 감지가 0이 아니고, 정합화하면 0이 된다
  (워크북은 임시 복사본에서만 만진다 — 실물 파일은 건드리지 않는다)

실행: venv312\\Scripts\\python.exe scratch\\test_block_realign.py
"""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

import candidate_tracker_auto_update as tracker  # noqa: E402

REAL_FILE = ROOT / "★조합생성_후보숫자_추적표" / "★후보숫자_추적표_표본vs최근50회.xlsx"
NUMS = list(range(1, 46))
TMP = Path(tempfile.mkdtemp(prefix="lotto_realign_test_"))


def setUpModule():
    tracker.LOG_FILE = TMP / "realign_test.log"


def tearDownModule():
    shutil.rmtree(TMP, ignore_errors=True)


def _draws(n: int = 60) -> list[tuple[int, list[int]]]:
    """결정적 합성 이력 — 회차 r의 번호는 seed 기반으로 고정."""
    import random

    rnd = random.Random(20260927)
    return [(r, sorted(rnd.sample(range(1, 46), 6))) for r in range(1000, 1000 + n)]


def _build_synthetic_workbook(path: Path, draws) -> None:
    """'밀리는 블록' 구조를 가진 최소 워크북:
       전체당첨내역(A=회차, B..H=번호+보너스) + 블록 시트(A수식, L=값, I..K=적중수)."""
    wb = openpyxl.Workbook()
    ws_all = wb.active
    ws_all.title = "전체당첨내역"
    ws_all.append(["회차", "1구", "2구", "3구", "4구", "5구", "6구", "보너스"])
    for r, ns in draws:
        ws_all.append([r] + ns + [7])

    ws = wb.create_sheet("3차필터(합성)")
    ws.cell(2, 11, "기준빈도(전체)")
    ws.cell(3, 11, "최근10회 빈도")
    latest = draws[-1][0]
    start = 8
    for i in range(6):                       # 6행짜리 블록
        r = start + i
        rnd = latest - i                     # 시작행 = 최신회차행(offset 0)
        ws.cell(r, 1, f"=MAX(전체당첨내역!$A$2:$A$100){'-' + str(i) if i else ''}")
        drawn = next(ns for rr, ns in draws if rr == rnd)
        for j, n in enumerate(drawn):        # B..G: 그 회차 당첨번호(실물과 같은 역할)
            ws.cell(r, 2 + j, n)
        ws.cell(r, 8, 7)                     # H: 보너스
        for c in range(12, 57):              # L..BD: 일부러 틀린 값(45..1)
            ws.cell(r, c, 46 - (c - 11))
        for c in (9, 10, 11):                # I..K: 일부러 틀린 적중수
            ws.cell(r, c, 0)
    wb.save(path)
    wb.close()


class PredictionDefinitionTests(unittest.TestCase):
    """A·B) 예측의 정의 — 미래를 보지 않고, 항상 순열이다."""

    def test_prediction_is_a_permutation(self):
        draws = _draws(60)
        for target in (1001, 1030, 1060):
            got = tracker._prediction_for(draws, target, None, 10)
            with self.subTest(target=target):
                self.assertIsNotNone(got)
                self.assertEqual(sorted(got), NUMS)

    def test_prediction_ignores_rounds_at_or_after_target(self):
        """X의 예측은 X-1까지만 봐야 한다 — X 이후 데이터를 지워도/더해도 같아야."""
        draws = _draws(60)
        target = 1040
        truncated = [(r, ns) for r, ns in draws if r < target]
        extended = list(draws) + [(1070, [1, 2, 3, 4, 5, 6]), (1071, [7, 8, 9, 10, 11, 12])]
        base = tracker._prediction_for(truncated, target, None, 10)
        self.assertEqual(tracker._prediction_for(draws, target, None, 10), base)
        self.assertEqual(tracker._prediction_for(extended, target, None, 10), base)

    def test_no_history_before_target_gives_none(self):
        draws = _draws(5)
        self.assertIsNone(tracker._prediction_for(draws, 1000, None, 10))

    def test_window_changes_ranking(self):
        """창을 바꾸면 순위가 달라질 수 있어야 한다 — 두 창이 늘 같은 값을 내면 창 처리가 죽은 것."""
        draws = _draws(60)
        self.assertNotEqual(tracker._prediction_for(draws, 1060, None, 5),
                            tracker._prediction_for(draws, 1060, None, 50))

    def test_hits_conservation(self):
        """적중수는 어떤 순위/당첨번호 조합에서도 합이 6 (전 회차 번호가 45칸 안에 다 있으므로)."""
        draws = _draws(30)
        for r, ns in draws[:5]:
            ranking = list(NUMS)
            got = tracker._hits(ranking, ns)
            with self.subTest(round=r):
                self.assertEqual(sum(got), 6)
                self.assertEqual(got, (sum(1 for d in ns if d <= 15),
                                       sum(1 for d in ns if 16 <= d <= 30),
                                       sum(1 for d in ns if d > 30)))


class SyntheticWorkbookTests(unittest.TestCase):
    """C·D·E) 합성 워크북에서 감지→정합화→재검증 전 과정."""

    def setUp(self):
        self.path = TMP / f"syn_{self._testMethodName}.xlsx"
        self.draws = _draws(60)
        _build_synthetic_workbook(self.path, self.draws)

    def _realign(self):
        wb = openpyxl.load_workbook(self.path, data_only=False)
        try:
            fixed = tracker.realign_sliding_blocks(wb, self.draws)
            wb.save(self.path)
        finally:
            wb.close()
        return fixed

    def test_detects_broken_block_then_heals(self):
        self.assertGreater(tracker._block_mismatch_rows(self.path), 0, "어긋난 입력을 못 잡음")
        fixed = self._realign()
        self.assertTrue(fixed, "정합화 대상 시트를 못 찾음")
        self.assertEqual(tracker._block_mismatch_rows(self.path), 0, "정합화 후에도 어긋남")

    def test_every_row_pairs_its_own_prediction_with_its_own_numbers(self):
        self._realign()
        wb = openpyxl.load_workbook(self.path, data_only=True)
        try:
            ws = wb["3차필터(합성)"]
            latest = self.draws[-1][0]
            for r in range(8, 14):
                rnd = latest - (r - 8)
                ranking = [ws.cell(r, c).value for c in range(12, 57)]
                got = [ws.cell(r, c).value for c in (9, 10, 11)]
                drawn = next(ns for rr, ns in self.draws if rr == rnd)
                with self.subTest(row=r, round=rnd):
                    self.assertEqual(sorted(ranking), NUMS)                       # B) 순열
                    self.assertEqual(got, list(tracker._hits(ranking, drawn)))    # C) 자기 회차끼리 짝
                    self.assertEqual(sum(got), 6)
                    self.assertEqual(ranking, tracker._prediction_for(self.draws, rnd, None, 10))
        finally:
            wb.close()

    def test_idempotent(self):
        self._realign()
        first = openpyxl.load_workbook(self.path, data_only=False)
        snap1 = [[first["3차필터(합성)"].cell(r, c).value for c in range(1, 57)] for r in range(8, 14)]
        first.close()
        fixed_again = self._realign()          # 두 번째 적용
        second = openpyxl.load_workbook(self.path, data_only=False)
        snap2 = [[second["3차필터(합성)"].cell(r, c).value for c in range(1, 57)] for r in range(8, 14)]
        second.close()
        self.assertEqual(snap1, snap2, "정합화가 멱등적이지 않음")
        self.assertTrue(fixed_again, "두 번째에도 정합화가 돌아야 한다(멱등하지만 대상은 잡힘)")

    def test_sheet_without_sliding_block_is_untouched(self):
        """A가 수식이 아니면(= 스스로 갱신되지 않는 시트) 건드리지 않는다."""
        wb = openpyxl.load_workbook(self.path)
        ws = wb.create_sheet("고정시트")
        ws.cell(2, 11, "기준빈도(전체)")
        ws.cell(3, 11, "최근10회 빈도")
        ws.cell(8, 1, 9999)                    # 값(수식 아님)
        for c in range(12, 57):
            ws.cell(8, c, c - 11)
        wb.save(self.path)
        wb.close()
        before = openpyxl.load_workbook(self.path, data_only=False)
        snap = [before["고정시트"].cell(8, c).value for c in range(1, 57)]
        before.close()
        self._realign()
        after = openpyxl.load_workbook(self.path, data_only=False)
        self.assertEqual([after["고정시트"].cell(8, c).value for c in range(1, 57)], snap)
        after.close()


class RealWorkbookTests(unittest.TestCase):
    """실물 워크북(임시 복사본)에서 검사 — 지금 파일이 정합 상태인지도 함께 확인."""

    @classmethod
    def setUpClass(cls):
        if not REAL_FILE.exists():
            raise unittest.SkipTest("실물 워크북 없음")
        cls.copy = TMP / "real_copy.xlsx"
        shutil.copy2(REAL_FILE, cls.copy)

    def test_real_file_is_consistent_and_realign_is_idempotent(self):
        self.assertEqual(tracker._block_mismatch_rows(self.copy), 0,
                         "실물 파일이 아직 어긋나 있다 — 정합화가 필요")
        wb = openpyxl.load_workbook(self.copy, data_only=False)
        try:
            draws = tracker._draws_from_sheet(wb["전체당첨내역"])
            snap = [[wb["3차필터(500회_후보)"].cell(r, c).value for c in range(1, 57)]
                    for r in range(7, 57)]
            tracker.realign_sliding_blocks(wb, draws)     # 다시 돌려도
            again = [[wb["3차필터(500회_후보)"].cell(r, c).value for c in range(1, 57)]
                     for r in range(7, 57)]
        finally:
            wb.close()
        self.assertEqual(snap, again, "실물 파일에서 재적용 결과가 달라짐(멱등성 위반)")


if __name__ == "__main__":
    unittest.main(verbosity=2)
