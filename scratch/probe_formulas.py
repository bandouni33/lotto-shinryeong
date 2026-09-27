# -*- coding: utf-8 -*-
"""두 윈도우비교 파일의 수식 구조 + 백업본 구조 + 전체당첨내역 꼬리 진단(읽기 전용)."""
from __future__ import annotations

from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TRACKER = ROOT / "★조합생성_후보숫자_추적표"
OUT = ROOT / "scratch" / "probe_formulas_out.txt"
TARGETS = {
    "파일1 기준빈도(전체)": TRACKER / "조합생성_후보숫자_추적표_전체표본_윈도우비교.xlsx",
    "파일2 기준빈도(최근500회)": TRACKER / "조합생성_후보숫자_추적표_최근500표본_윈도우비교.xlsx",
}
BACKUPS = {
    "백업 전체표본(09-13)": TRACKER / "_backup_20260927" / "조합생성_후보숫자_추적표_전체표본_윈도우비교.xlsx",
    "백업 최근500표본(09-15)": TRACKER / "_backup_20260927" / "조합생성_후보숫자_추적표_최근500표본_윈도우비교.xlsx",
}

lines: list[str] = []


def w(s: str = "") -> None:
    lines.append(s)


def txt(v):
    if hasattr(v, "text"):
        return f"ArrayFormula(ref={v.ref}, text={v.text})"
    return repr(v)


w("== 백업본(09-27 10:57 이전 상태) 시트 구조 ==")
for label, p in BACKUPS.items():
    if not p.exists():
        w(f"  {label}: 없음")
        continue
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    try:
        w(f"  {label}: sheets={wb.sheetnames}")
        for sn in wb.sheetnames:
            if sn.startswith("3차필터"):
                w(f"     {sn}: K2={wb[sn].cell(2, 11).value!r} K3={wb[sn].cell(3, 11).value!r}")
        ws = wb["전체당첨내역"]
        rounds = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)
                  if isinstance(ws.cell(r, 1).value, int)]
        w(f"     전체당첨내역 max_row={ws.max_row} latest={max(rounds)}")
    finally:
        wb.close()

for label, path in TARGETS.items():
    w(f"\n{'=' * 70}\n{label} -> {path.name}")
    wbf = openpyxl.load_workbook(path, data_only=False)   # 수식
    wbv = openpyxl.load_workbook(path, data_only=True)    # 캐시값
    try:
        w("  -- 전체당첨내역 꼬리(마지막 3행, 값) --")
        wsv = wbv["전체당첨내역"]
        for r in range(wsv.max_row - 2, wsv.max_row + 1):
            w(f"     row{r}: {[wsv.cell(r, c).value for c in range(1, 9)]}")
        w(f"     max_row={wsv.max_row} max_column={wsv.max_column}")
        for sn in [s for s in wbf.sheetnames if s.startswith("3차필터")]:
            wsf, wsv = wbf[sn], wbv[sn]
            w(f"  -- {sn} (rows={wsf.max_row})")
            w(f"     K2={wsf.cell(2, 11).value!r} K3={wsf.cell(3, 11).value!r} K4={wsf.cell(4, 11).value!r}")
            for coord in ("A7", "A8", "B7", "G7", "I7", "L2", "L3", "L4", "L7", "M7", "BD7",
                          "L8", "A6", "E6", "A56" if wsf.max_row == 56 else f"A{wsf.max_row}"):
                cell_f = wsf[coord]
                cell_v = wsv[coord]
                w(f"     {coord}: formula={txt(cell_f.value)} | cached={cell_v.value!r}")
            # 블록 L열이 수식인지 값인지
            lvals = [wsf.cell(7, c).value for c in range(12, 57)]
            kinds = {"formula": sum(1 for v in lvals if isinstance(v, str) or hasattr(v, "text")),
                     "int": sum(1 for v in lvals if isinstance(v, int))}
            w(f"     7행 L:BD 종류 = {kinds}")
            lvals8 = [wsf.cell(8, c).value for c in range(12, 57)]
            w(f"     8행 L:BD 종류 = {{'formula': {sum(1 for v in lvals8 if isinstance(v, str) or hasattr(v, 'text'))}, "
              f"'int': {sum(1 for v in lvals8 if isinstance(v, int))}}}")
            w(f"     1행 L:BD = {[wsf.cell(1, c).value for c in range(12, 57)][:10]}...")
            datas = wsf.max_row - 7 + 1
            w(f"     블록행수(7..max) = {datas}")
    finally:
        wbf.close()
        wbv.close()

OUT.write_text("\n".join(lines), encoding="utf-8")
print(f"written {OUT} ({len(lines)} lines)")
