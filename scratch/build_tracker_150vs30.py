# -*- coding: utf-8 -*-
"""2026-10-11 최근150vs30 추적표 생성 기록(최근500표본 복사 → 3차필터(150회_후보) 1개만 남김 → 2행 기준 최근150·3행 최근30·5행 최근30회차 집계 → 과거 149회차 재계산).
실행: python scratch/build_tracker_150vs30.py <최근500표본.xlsx> <새파일.xlsx>  (검증: LibreOffice 재계산 7행 = 파이썬 폴백, 5행 = 직접 집계, 오류값 0, 1회 전진 모의 통과)"""
import openpyxl, re, sys
from openpyxl.worksheet.formula import ArrayFormula
SRC, DST = sys.argv[1], sys.argv[2]
BASE, REC, HIT_N = 150, 30, 30
wb = openpyxl.load_workbook(SRC)
print("defined names:", list(wb.defined_names.keys()) if hasattr(wb.defined_names,'keys') else wb.defined_names)
keep = "3차필터(150회_후보)"
for sn in list(wb.sheetnames):
    if sn not in ("전체당첨내역", keep):
        del wb[sn]
ws = wb[keep]
ws.title = f"3차필터({REC}회_후보)"
# 2행: 기준 최근150 / 3행: 최근30
for c in range(12, 57):
    f2 = ws.cell(2, c).value; f3 = ws.cell(3, c).value
    assert "-499))" in f2 and "-149))" in f3, (c, f2[:80], f3[:80])
    ws.cell(2, c).value = f2.replace("-499))", f"-{BASE-1}))")
    ws.cell(3, c).value = f3.replace("-149))", f"-{REC-1}))")
    f5 = ws.cell(5, c).value
    assert "$7:" in f5 and "$156" in f5
    ws.cell(5, c).value = f5.replace("$7:", "$8:").replace("$156", f"${7+HIT_N}")
ws.cell(2, 11).value = f"기준빈도(최근{BASE}회)"
ws.cell(3, 11).value = f"최근{REC}회 빈도"
ws.cell(5, 11).value = f"적중회수(1위 기준, 최근 {HIT_N}회차)"
# 과거 확정행 재계산
draws = {}
for row in wb["전체당첨내역"].iter_rows(min_row=2, values_only=True):
    if row[0] is None: continue
    draws[int(row[0])] = [int(x) for x in row[1:7]]
rounds = sorted(draws)
def counts(anchor, window):
    lo = anchor - (window - 1)
    c = {n: 0 for n in range(1, 46)}
    for r in rounds:
        if lo <= r <= anchor:
            for n in draws[r]: c[n] += 1
    return c
def rank(c):
    o = sorted(range(1, 46), key=lambda n: (-c[n], n)); return {n: i + 1 for i, n in enumerate(o)}
def forecast(anchor):
    rb, rr = rank(counts(anchor, BASE)), rank(counts(anchor, REC))
    return sorted(range(1, 46), key=lambda n: (-(rr[n] - rb[n]), n))
n_rows = 0
for r in range(8, ws.max_row + 1):
    rnd = ws.cell(r, 1).value
    if not isinstance(rnd, int): continue
    actual = [ws.cell(r, c).value for c in range(2, 8)]
    assert sorted(actual) == sorted(draws[rnd]), (rnd, actual, draws[rnd])
    fc = forecast(rnd - 1)
    for i, c in enumerate(range(12, 57)): ws.cell(r, c).value = fc[i]
    a = set(actual)
    ws.cell(r, 9).value = len(a & set(fc[:15])); ws.cell(r, 10).value = len(a & set(fc[15:30])); ws.cell(r, 11).value = len(a & set(fc[30:]))
    n_rows += 1
print("rows recomputed", n_rows, "range", ws.cell(8,1).value, "->", ws.cell(ws.max_row,1).value, "pending", ws.cell(7,1).value)
wb.active = 1
wb.calculation.fullCalcOnLoad = True  # 캐시값이 없으니 엑셀이 열 때 전부 다시 계산
wb.save(DST)
