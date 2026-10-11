# -*- coding: utf-8 -*-
"""2026-10-11 최근30회 추적표 생성 기록: 최근500표본 복사 → 3차필터 100·150·200·250·300 시트만 남김 → 2행 기준=최근NNN·3행=최근30·5행=최근30회차(8~37행) 집계 → 과거 회차 재계산(앵커 R-1), 시트명 3차필터(NNNvs30회_후보).
실행: python scratch/build_tracker_30vs.py <최근500표본.xlsx> <새파일.xlsx>  검증: LibreOffice 재계산 7행=파이썬 폴백(5시트), 5행=직접 집계, 오류값 0, 1246 모의 전진 통과."""
import openpyxl, re, sys
SRC, DST = sys.argv[1], sys.argv[2]
REC, HIT_N = 30, 30
BASES = (100, 150, 200, 250, 300)
wb = openpyxl.load_workbook(SRC)
for sn in list(wb.sheetnames):
    if sn != "전체당첨내역" and sn not in [f"3차필터({b}회_후보)" for b in BASES]:
        del wb[sn]
draws = {}
for row in wb["전체당첨내역"].iter_rows(min_row=2, values_only=True):
    if row[0] is None: continue
    draws[int(row[0])] = [int(x) for x in row[1:7]]
rounds = sorted(draws)
def counts(anchor, window):
    lo = anchor - (window - 1); c = {n: 0 for n in range(1, 46)}
    for r in rounds:
        if lo <= r <= anchor:
            for n in draws[r]: c[n] += 1
    return c
def rank(c):
    o = sorted(range(1, 46), key=lambda n: (-c[n], n)); return {n: i + 1 for i, n in enumerate(o)}
def forecast(anchor, base):
    rb, rr = rank(counts(anchor, base)), rank(counts(anchor, REC))
    return sorted(range(1, 46), key=lambda n: (-(rr[n] - rb[n]), n))
for b in BASES:
    ws = wb[f"3차필터({b}회_후보)"]
    for c in range(12, 57):
        f2, f3, f5 = ws.cell(2, c).value, ws.cell(3, c).value, ws.cell(5, c).value
        assert "-499))" in f2 and f"-{b-1}))" in f3, (b, c)
        ws.cell(2, c).value = f2.replace("-499))", f"-{b-1}))")
        ws.cell(3, c).value = f3.replace(f"-{b-1}))", f"-{REC-1}))")
        f5n = re.sub(r"\$7:", "$8:", f5)
        f5n = re.sub(r"\$(\d+)", lambda m: f"${7+HIT_N}" if int(m.group(1)) >= 50 else m.group(0), f5n)
        assert f5n.count(f"${7+HIT_N}") == 13 and f5n.count("$8:") == 13, f5n[:160]
        ws.cell(5, c).value = f5n
    ws.cell(2, 11).value = f"기준빈도(최근{b}회)"
    ws.cell(3, 11).value = f"최근{REC}회 빈도"
    ws.cell(5, 11).value = f"적중회수(1위 기준, 최근 {HIT_N}회차)"
    n = 0
    for r in range(8, ws.max_row + 1):
        rnd = ws.cell(r, 1).value
        if not isinstance(rnd, int): continue
        actual = [ws.cell(r, c).value for c in range(2, 8)]
        assert sorted(actual) == sorted(draws[rnd]), (b, rnd)
        fc = forecast(rnd - 1, b)
        for i, c in enumerate(range(12, 57)): ws.cell(r, c).value = fc[i]
        a = set(actual)
        ws.cell(r, 9).value = len(a & set(fc[:15])); ws.cell(r, 10).value = len(a & set(fc[15:30])); ws.cell(r, 11).value = len(a & set(fc[30:]))
        n += 1
    ws.title = f"3차필터({b}vs{REC}회_후보)"
    print(ws.title, "rows", n, ws.cell(8, 1).value, "->", ws.cell(ws.max_row, 1).value, "pending", ws.cell(7, 1).value)
wb.active = 1
wb.calculation.fullCalcOnLoad = True
wb.save(DST)
print(wb.sheetnames)
