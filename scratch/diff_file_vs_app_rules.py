# -*- coding: utf-8 -*-
"""파일 1차 규칙표 vs 앱 규칙(DB/JSON) 1:1 차이 덤프 (읽기 전용).

배경(2026-09-27 실측): 200회 AUTO 규칙을 앱에 추가한 뒤에도 앱 파이프라인의
1차+이격 통과수가 2,594,759로, 파일 규칙으로 계산한 2,461,980과 5.4% 어긋났다.
AUTO(4개)와 이격수(48개)는 파일과 같은 것을 확인했으므로, 남은 차이는 **고정 규칙
378개 내용**에서 온다. 어느 행이 어떻게 다른지 그대로 뽑아 보여준다.

실행: venv312\\Scripts\\python.exe scratch\\diff_file_vs_app_rules.py
"""
from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

import combo_filter_v2 as cf  # noqa: E402

from env_loader import load_dotenv_file  # noqa: E402

load_dotenv_file()

SRC = ROOT / "★조합생성_후보숫자_추적표" / "조합생성_후보숫자_추적표_샘플.xlsx"
OUT = ROOT / "scratch" / "diff_file_vs_app_rules_out.txt"
R: list[str] = []


def w(s: str = "") -> None:
    R.append(str(s))
    print(s, flush=True)


def tgt(v) -> list[int]:
    return sorted({int(x) for x in re.findall(r"\d+", str(v))}) if v is not None else []


wb = openpyxl.load_workbook(SRC, data_only=True)
try:
    ws = wb["1차필터(7기본필터)"]
    file_rules = {}
    file_auto = {}
    for r in range(5, 1504):
        j, k, lm = ws.cell(r, 10).value, ws.cell(r, 11).value, ws.cell(r, 12).value
        if j is None or not isinstance(k, (int, float)) or not isinstance(lm, (int, float)):
            continue
        if str(j).strip().upper() == "AUTO":
            file_auto[r] = {"name": ws.cell(r, 8).value, "min": int(k), "max": int(lm)}
        else:
            t = tgt(j)
            if t:
                file_rules[r] = {"targets": t, "min": int(k), "max": int(lm),
                                 "name": ws.cell(r, 8).value}
finally:
    wb.close()

static, auto, gap = cf._load_rules()
app_rules = {int(r.get("row", -1)): {"targets": sorted({int(x) for x in r["targets"]}),
                                     "min": int(r["min"]), "max": int(r["max"]),
                                     "name": r.get("name")} for r in static}
app_auto = {int(r.get("row", -1)): {"name": r.get("name"), "min": int(r["min"]),
                                    "max": int(r["max"])} for r in auto}
app_rows = {int(r.get("row", -1)) for r in static}

w(f"파일 고정 규칙 {len(file_rules)}개 / 앱 고정 규칙 {len(static)}개")
w(f"파일 AUTO {len(file_auto)}개 {[(k, v['name']) for k, v in file_auto.items()]}")
w(f"앱   AUTO {len(app_auto)}개 {[(k, v['name']) for k, v in app_auto.items()]}")

only_file = sorted(set(file_rules) - app_rows)
only_app = sorted(app_rows - set(file_rules))
w(f"\n파일에만 있는 행 {len(only_file)}개: {only_file[:30]}")
w(f"앱에만 있는 행 {len(only_app)}개: {only_app[:30]}")

diffs = []
for row in sorted(set(file_rules) & app_rows):
    f, a = file_rules[row], app_rules[row]
    if (f["targets"], f["min"], f["max"]) != (a["targets"], a["min"], a["max"]):
        diffs.append((row, f, a))
w(f"\n내용이 다른 행 {len(diffs)}개 (행번호 · 파일 개수/범위 · 앱 개수/범위 · 대칭차)")
for row, f, a in diffs[:40]:
    sym = sorted(set(f["targets"]) ^ set(a["targets"]))
    w(f"  row{row}: 파일 {len(f['targets'])}개 {f['min']}~{f['max']} / "
      f"앱 {len(a['targets'])}개 {a['min']}~{a['max']} / 차이 {len(sym)}개 {sym[:8]}")
if len(diffs) > 40:
    w(f"  … 총 {len(diffs)}개 중 40개만 표시")

auto_diff = [(r, file_auto[r], app_auto.get(r)) for r in sorted(set(file_auto) | set(app_auto))
             if file_auto.get(r) != app_auto.get(r)]
w(f"\nAUTO 차이 {len(auto_diff)}개:")
for row, f, a in auto_diff:
    w(f"  row{row}: 파일 {f} / 앱 {a}")

w("\n== 요약 ==")
w(f"  · 고정 규칙: 파일 {len(file_rules)} vs 앱 {len(static)} — 내용 불일치 {len(diffs)}개, "
  f"한쪽에만 있는 행 {len(only_file) + len(only_app)}개, AUTO 차이 {len(auto_diff)}개")
if diffs or only_file or only_app or auto_diff:
    w("  · 남은 차이가 있으면 그 차이가 1차+이격 통과수 차이의 원인이다")
else:
    w("  · 규칙 완전 일치: 파일 규칙표와 앱 규칙(DB)이 고정 378 · AUTO 4 · 이격수 48 모두 동일")
    w("    (2026-09-27 row19 max 3→2 수정으로 통일 완료)")

OUT.write_text("\n".join(R), encoding="utf-8")
print(f"written {OUT}")
os._exit(0)
