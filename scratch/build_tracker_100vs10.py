# -*- coding: utf-8 -*-
"""2026-10-11 '100vs10 비교표' 생성 기록(사용자 지시): 최근30회 윈도우비교 파일 복사 → 5개 시트
3행 비교창 최근30회 → 최근10회, 5행 적중회수 최근 30회차(8~37행) → 최근 10회차(8~17행), 시트명
3차필터(NNNvs10회_후보) → 과거 회차 예측(L:BD)·적중(I/J/K) 재계산(앵커 R-1). 서식·나머지는 그대로.
실행: python scratch/build_tracker_100vs10.py <최근30회.xlsx> <새파일.xlsx>"""
import openpyxl, sys
SRC, DST = sys.argv[1], sys.argv[2]
OLD_REC, REC, HIT_N = 30, 10, 10
BASES = (100, 150, 200, 250, 300)
wb = openpyxl.load_workbook(SRC)
draws = {}
for row in wb["전체당첨내역"].iter_rows(min_row=2, values_only=True):
    if row[0] is None: continue
    draws[int(row[0])] = [int(x) for x in row[1:7]]
rounds = sorted(draws)
def counts(anchor, window):
    lo = anchor - (window - 1); c = {n: 0 for n in range(1, 46)}
    for r in rounds:
        if lo <= r <= anchor:
            for n in draws[r]: c[n] += 1
    return c
def rank(c):
    o = sorted(range(1, 46), key=lambda n: (-c[n], n)); return {n: i + 1 for i, n in enumerate(o)}
def forecast(anchor, base):
    rb, rr = rank(counts(anchor, base)), rank(counts(anchor, REC))
    return sorted(range(1, 46), key=lambda n: (-(rr[n] - rb[n]), n))
for b in BASES:
    ws = wb[f"3차필터({b}vs{OLD_REC}회_후보)"]
    for c in range(12, 57):
        f3, f5 = ws.cell(3, c).value, ws.cell(5, c).value
        assert f"-{OLD_REC-1}))" in f3 and f5.count("$37") == 13 and f5.count("$8:") == 13, (b, c)
        ws.cell(3, c).value = f3.replace(f"-{OLD_REC-1}))", f"-{REC-1}))")
        ws.cell(5, c).value = f5.replace("$37", f"${7+HIT_N}")
    assert ws.cell(2, 11).value == f"기준빈도(최근{b}회)"
    ws.cell(3, 11).value = f"최근{REC}회 빈도"
    ws.cell(5, 11).value = f"적중회수(1위 기준, 최근 {HIT_N}회차)"
    n = 0
    for r in range(8, ws.max_row + 1):
        rnd = ws.cell(r, 1).value
        if not isinstance(rnd, int): continue
        actual = [ws.cell(r, c).value for c in range(2, 8)]
        assert sorted(actual) == sorted(draws[rnd]), (b, rnd)
        fc = forecast(rnd - 1, b)
        for i, c in enumerate(range(12, 57)): ws.cell(r, c).value = fc[i]
        a = set(actual)
        ws.cell(r, 9).value = len(a & set(fc[:15])); ws.cell(r, 10).value = len(a & set(fc[15:30])); ws.cell(r, 11).value = len(a & set(fc[30:]))
        n += 1
    ws.title = f"3차필터({b}vs{REC}회_후보)"
    print(ws.title, "rows", n, ws.cell(8, 1).value, "->", ws.cell(ws.max_row, 1).value, "pending", ws.cell(7, 1).value)
# 시트 이름을 바꿨으니 다른 시트(Sheet1 등)의 수식 속 시트 참조도 같이 바꾼다(openpyxl은 자동으로 안 바꿈).
from openpyxl.worksheet.formula import ArrayFormula
fixed = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            t = v.text if isinstance(v, ArrayFormula) else v
            if isinstance(t, str) and t.startswith("=") and f"vs{OLD_REC}회_후보" in t:
                nt = t.replace(f"vs{OLD_REC}회_후보", f"vs{REC}회_후보")
                cell.value = ArrayFormula(ref=v.ref, text=nt) if isinstance(v, ArrayFormula) else nt
                fixed += 1
print("시트 참조 수정", fixed)
wb.calculation.fullCalcOnLoad = True
wb.save(DST)
print(wb.sheetnames)
