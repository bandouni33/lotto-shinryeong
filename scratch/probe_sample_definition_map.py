# -*- coding: utf-8 -*-
"""사용자가 준 정의와 파일 구조를 맞대기 위한 정밀 읽기(읽기 전용).

확인:
  · 1차필터(7기본필터) 규칙표(H/J/K/L) 전체 + 특정 회차 열의 판정/헬퍼행
  · 2차필터(5이격수) 규칙표 + 회차 열
  · 1차/2차/4차 추적결과의 머리글(A~S)과 최신 회차행 값 → H/I 가 무엇인지
  · 후보시트(68열)의 2~5행 × H~BK (사용자 설명: S2:BK2 역대 / S3:BK3 최근100 / S4:BK4 오차 / S5:BK5 적중회수)

실행: venv312\\Scripts\\python.exe scratch\\probe_sample_definition_map.py
"""
from __future__ import annotations

import os
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
SRC = ROOT / "★조합생성_후보숫자_추적표" / "조합생성_후보숫자_추적표_샘플.xlsx"
OUT = ROOT / "scratch" / "probe_sample_definition_map_out.txt"
L: list[str] = []


def w(s: str = "") -> None:
    L.append(s)


def cell(ws, r, c):
    v = ws.cell(r, c).value
    return v.text if hasattr(v, "text") else v


def short(v, n=70):
    s = repr(v)
    return s if len(s) <= n else s[:n] + "…"


wb = openpyxl.load_workbook(SRC, data_only=True)
try:
    # ── 1차필터 규칙표: 값 보기
    ws1 = wb["1차필터(7기본필터)"]
    w("=== 1차필터(7기본필터) 규칙표 (H/J/K/L), 4~30행")
    for r in range(4, 31):
        vals = [cell(ws1, r, c) for c in (8, 10, 11, 12)]
        if any(v is not None for v in vals):
            w(f"  row{r}: H={short(vals[0])} J={short(vals[1])} K={short(vals[2])} L={short(vals[3])}")
    # 회차 열 찾기(4행 라벨)
    cols = {cell(ws1, 4, c): c for c in range(1, ws1.max_column + 1)}
    w(f"  4행 라벨(회차→열): {cols}")
    c1241 = cols.get(1241)
    if c1241:
        w(f"  -- 1241 열({openpyxl.utils.get_column_letter(c1241)}) 1~20행 값:")
        for r in range(1, 21):
            w(f"     r{r}={short(cell(ws1, r, c1241))}")
        w(f"  -- 같은 열 1495~1516행 값:")
        for r in range(1495, 1517):
            w(f"     r{r}={short(cell(ws1, r, c1241))}")
        w(f"  -- 같은 열 1~20행 수식:")
        wf1 = openpyxl.load_workbook(SRC, data_only=False, read_only=True)
        try:
            ws1f = wf1["1차필터(7기본필터)"]
            for r in range(1, 21):
                row = next(ws1f.iter_rows(min_row=r, max_row=r, min_col=c1241, max_col=c1241,
                                          values_only=True))
                v = row[0]
                t = v.text if hasattr(v, "text") else v
                if t is not None:
                    w(f"     r{r}={short(t, 160)}")
        finally:
            wf1.close()

    # ── 2차필터 규칙표
    ws2 = wb["2차필터(5이격수)"]
    w("\n=== 2차필터(5이격수) 규칙표 (H/J/K/L), 4~30행")
    for r in range(4, 31):
        vals = [cell(ws2, r, c) for c in (8, 10, 11, 12)]
        if any(v is not None for v in vals):
            w(f"  row{r}: H={short(vals[0])} J={short(vals[1])} K={short(vals[2])} L={short(vals[3])}")
    w(f"  4행 라벨(회차→열): {[ (k,v) for k,v in {cell(ws2,4,c): c for c in range(1, ws2.max_column+1)}.items() if isinstance(k,int)]}")

    # ── 추적결과 머리글·최신행
    for sn in ("1차추적결과", "2차추적결과", "4차필터"):
        ws = wb[sn]
        w(f"\n=== {sn}: 머리글 1~2행 + 3행")
        w("  row1: " + ", ".join(f"{openpyxl.utils.get_column_letter(c)}={short(cell(ws,1,c),30)}"
                                 for c in range(1, 20) if cell(ws, 1, c) is not None))
        w("  row2: " + ", ".join(f"{openpyxl.utils.get_column_letter(c)}={short(cell(ws,2,c),30)}"
                                 for c in range(1, 20) if cell(ws, 2, c) is not None))
        w("  row3: " + ", ".join(f"{openpyxl.utils.get_column_letter(c)}={short(cell(ws,3,c),40)}"
                                 for c in range(1, 20) if cell(ws, 3, c) is not None))

    # ── 후보시트 헤더 구조
    for sn in [s for s in wb.sheetnames if "출현빈도순" in s]:
        ws = wb[sn]
        w(f"\n=== {sn}: rows={ws.max_row} cols={ws.max_column}")
        for r in range(1, 7):
            parts = []
            for c in range(1, ws.max_column + 1):
                v = cell(ws, r, c)
                if v is not None:
                    parts.append(f"{openpyxl.utils.get_column_letter(c)}{r}={short(v,40)}")
            if parts:
                w(f"  row{r}: " + " | ".join(parts[:14]) + (" …" if len(parts) > 14 else ""))
finally:
    wb.close()

OUT.write_text("\n".join(L), encoding="utf-8")
print(f"written {OUT} ({len(L)} lines)")
