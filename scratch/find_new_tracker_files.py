# -*- coding: utf-8 -*-
"""K2='기준빈도(전체)' / '기준빈도(최근500회)' + 3차필터 6시트 + 전체당첨내역 1244 구조의
엑셀 파일을 찾아 UTF-8 리포트로 덤프한다(읽기 전용). 바탕화면 전체를 훑는다.
"""
from __future__ import annotations

import datetime
import os
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "scratch" / "find_new_tracker_files_out.txt"
DESKTOP = Path(r"C:\Users\PC\Desktop")

lines: list[str] = []


def w(s: str = "") -> None:
    lines.append(s)


w(f"scan start {datetime.datetime.now():%Y-%m-%d %H:%M:%S}")

# 1) 추적표 폴더 원시 목록(숨김/임시 포함)
TRACKER = ROOT / "★조합생성_후보숫자_추적표"
w(f"\n== raw listdir of {TRACKER} ==")
try:
    for e in sorted(os.listdir(TRACKER)):
        p = TRACKER / e
        w(f"  {e!r} dir={p.is_dir()} size={p.stat().st_size if p.is_file() else '-'} "
          f"mtime={datetime.datetime.fromtimestamp(p.stat().st_mtime):%Y-%m-%d %H:%M}")
except Exception as e:  # noqa: BLE001
    w(f"  !! {e}")

# 2) 바탕화면 전체 xlsx 훑기(최근 7일 이내 수정된 것 위주로 구조 판정)
w(f"\n== xlsx scan under {DESKTOP} ==")
cutoff = datetime.datetime.now() - datetime.timedelta(days=7)
hits: list[str] = []
for dirpath, dirnames, filenames in os.walk(DESKTOP):
    dirnames[:] = [d for d in dirnames if d not in {"node_modules", ".git", "venv312", ".venv"}]
    for fn in filenames:
        if not fn.lower().endswith(".xlsx") or fn.startswith("~$"):
            continue
        p = Path(dirpath) / fn
        try:
            mt = datetime.datetime.fromtimestamp(p.stat().st_mtime)
        except OSError:
            continue
        if mt < cutoff:
            continue
        try:
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            w(f"  !! open failed {p}: {e}")
            continue
        try:
            sn = wb.sheetnames
            filters = [s for s in sn if s.startswith("3차필터")]
            latest = None
            if "전체당첨내역" in sn:
                ws = wb["전체당첨내역"]
                rounds = [ws.cell(r, 1).value for r in range(2, min(ws.max_row, 5000) + 1)
                          if isinstance(ws.cell(r, 1).value, int)]
                latest = max(rounds) if rounds else None
            k2s = []
            for s in filters[:2]:
                if s in sn:
                    k2s.append(f"{s}:K2={wb[s].cell(2, 11).value!r}")
            w(f"  {p}")
            w(f"     mtime={mt:%Y-%m-%d %H:%M} sheets={len(sn)} filters={len(filters)} "
              f"latest_round={latest}")
            w(f"     {k2s}")
            if filters:
                hits.append(f"{p} | filters={len(filters)} | latest={latest} | {k2s}")
        finally:
            wb.close()

w("\n== files with 3차필터 sheets ==")
for h in hits:
    w("  " + h)

# 3) DB 최신 회차(읽기 전용) — 파일이 뒤처졌는지 판단용
w("\n== DB draw_results latest round ==")
try:
    import sys
    sys.path.insert(0, str(ROOT))
    import draw_results_db
    rows = draw_results_db.get_all_draw_results()
    rounds = sorted(int(r["draw_round"]) for r in rows)
    w(f"  n={len(rounds)} min={rounds[0] if rounds else None} max={rounds[-1] if rounds else None}")
    if rounds:
        last = next(r for r in rows if int(r["draw_round"]) == rounds[-1])
        w(f"  last={last}")
except Exception as e:  # noqa: BLE001
    w(f"  !! DB read failed: {type(e).__name__}: {e}")

OUT.write_text("\n".join(lines), encoding="utf-8")
print(f"written {OUT} ({len(lines)} lines)")
