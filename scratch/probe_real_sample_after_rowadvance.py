# -*- coding: utf-8 -*-
"""실물 샘플 파일 검증 — 행 전진·대조 시트가 실제로 어떻게 들어갔는지 (읽기 전용).

검사:
  1) 백업(scratch/_backup_샘플_before_rowadvance.xlsx)과 시트별 내용을 대조해
     **바뀐 시트가 행 전진 대상 3개 + 새 대조 시트뿐인지** 확인.
  2) 추적결과 창 라벨이 연속·내림차순으로 1243까지 왔는지.
  3) 대조 시트 내용(10열 전부)을 그대로 출력.
  4) 오류값 0 (weekly_lotto_file_update.scan_errors 재사용).

실행: venv312\\Scripts\\python.exe scratch\\probe_real_sample_after_rowadvance.py
"""
from __future__ import annotations

import hashlib
import os
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

import weekly_lotto_file_update as wk  # noqa: E402

REAL = ROOT / "★조합생성_후보숫자_추적표" / "조합생성_후보숫자_추적표_샘플.xlsx"
BACKUP = ROOT / "scratch" / "_backup_샘플_before_rowadvance.xlsx"
OUT = ROOT / "scratch" / "probe_real_sample_after_rowadvance_out.txt"
R: list[str] = []
FAILS: list[str] = []


def w(s: str = "") -> None:
    R.append(str(s))
    try:
        print(s, flush=True)
    except UnicodeEncodeError:
        print(str(s).encode("cp949", errors="replace").decode("cp949", errors="replace"),
              flush=True)


def ok(cond: bool, msg: str) -> None:
    if cond:
        w(f"  ok   {msg}")
    else:
        FAILS.append(msg)
        w(f"  FAIL {msg}")


def fx(cell):
    v = cell.value
    return v.text if hasattr(v, "text") else v


def sheet_sig(ws) -> str:
    h = hashlib.sha256()
    for row in ws.iter_rows():
        for c in row:
            h.update(f"{c.coordinate}={fx(c)!r}|".encode())
    return h.hexdigest()


w(f"실물 {REAL.name}  sha(앞12)={hashlib.sha256(REAL.read_bytes()).hexdigest()[:12]}")
w(f"백업 {BACKUP.name}  sha(앞12)={hashlib.sha256(BACKUP.read_bytes()).hexdigest()[:12]}")

wb_new = openpyxl.load_workbook(REAL, data_only=False)
wb_old = openpyxl.load_workbook(BACKUP, data_only=False)
try:
    w(f"\n== 시트 목록 ==")
    w(f"  전: {list(wb_old.sheetnames)}")
    w(f"  후: {list(wb_new.sheetnames)}")
    ok(list(wb_new.sheetnames) == list(wb_old.sheetnames) + [wk.DAEJO_SHEET],
       f"시트 = 기존 + {wk.DAEJO_SHEET} 하나")

    w("\n== 내용이 바뀐 시트 ==")
    changed = []
    for sn in wb_old.sheetnames:
        if sheet_sig(wb_old[sn]) != sheet_sig(wb_new[sn]):
            changed.append(sn)
    w(f"  {changed}")
    expect = set(wk.ROW_ROUND_SHEETS)
    ok(set(changed) == expect,
       f"바뀐 시트가 행 전진 대상 3개뿐 (기대 {sorted(expect)}, 실제 {sorted(changed)})")

    w("\n== 추적결과 창 ==")
    for sn in wk.ROW_ROUND_SHEETS:
        ws = wb_new[sn]
        labels = [ws.cell(r, 1).value for r in range(wk.ROW_FIRST, wk.ROW_LAST + 1)]
        ints = [x for x in labels if isinstance(x, int)]
        head = [ws.cell(r, 1).value for r in range(wk.ROW_FIRST, wk.ROW_FIRST + 3)]
        w(f"  [{sn}] {ints[0]}..{ints[-1]} ({len(ints)}개) · 앞 3행 라벨 {head} · "
          f"max_row={ws.max_row}")
        ok(len(ints) == wk.ROW_WINDOW and ints[0] == 1243 and ints[-1] == 1142,
           f"{sn}: 창 1243..1142 (102행)")

    w(f"\n== 대조 시트 '{wk.DAEJO_SHEET}' 내용 ==")
    ws = wb_new[wk.DAEJO_SHEET]
    heads = [ws.cell(1, c).value for c in range(1, len(wk.DAEJO_HEADERS) + 1)]
    w(f"  헤더: {heads}")
    ok(tuple(heads) == wk.DAEJO_HEADERS, "헤더가 코드 상수와 같다")
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, len(wk.DAEJO_HEADERS) + 1)]
        if vals[0] is None:
            continue
        w(f"  {vals}")
        ok(isinstance(vals[0], int) and vals[0] < wk.RULE_VINTAGE_ROUND
           and vals[7] == "이력(규칙 불일치)",
           f"{vals[0]}회차: 규칙 빈티지 이전이므로 판정 '{vals[7]}'")
    ok(ws.max_row == 3, f"대조 행 수 {ws.max_row - 1}개 (1242·1243 두 회차)")
finally:
    wb_new.close()
    wb_old.close()

errs = wk.scan_errors(REAL)
ok(not errs, f"오류값 0 (발견 {len(errs)} {errs[:5]})")

w(f"\n단언 실패 {len(FAILS)}건" + ("" if not FAILS else ": " + " | ".join(FAILS[:5])))
OUT.write_text("\n".join(R), encoding="utf-8")
print(f"written {OUT}")
os._exit(1 if FAILS else 0)
