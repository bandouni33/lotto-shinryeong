# -*- coding: utf-8 -*-
"""두 윈도우비교 파일의 '순번 행' 수식 패턴 확인(읽기 전용, read_only 스트리밍).

사용: python scratch/probe_formula_pattern.py 1   (1=전체표본, 2=최근500표본)
결과는 scratch/probe_formula_pattern_out<N>.txt 로 UTF-8 저장.
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TRACKER = ROOT / "★조합생성_후보숫자_추적표"
TARGETS = {
    "1": ("파일1 K2=기준빈도(전체)", TRACKER / "조합생성_후보숫자_추적표_전체표본_윈도우비교.xlsx"),
    "2": ("파일2 K2=기준빈도(최근500회)", TRACKER / "조합생성_후보숫자_추적표_최근500표본_윈도우비교.xlsx"),
}

key = sys.argv[1] if len(sys.argv) > 1 else "1"
label, path = TARGETS[key]
OUT = ROOT / "scratch" / f"probe_formula_pattern_out{key}.txt"
lines: list[str] = []


def w(s: str = "") -> None:
    lines.append(s)


def short(v, n: int = 150) -> str:
    s = repr(v)
    return s if len(s) <= n else s[:n] + "..."


def row_of(ws, r: int) -> list:
    for row in ws.iter_rows(min_row=r, max_row=r, values_only=True):
        return list(row)
    return []


w(f"{label}: {path.name}")
w(f"size={path.stat().st_size} mtime={__import__('datetime').datetime.fromtimestamp(path.stat().st_mtime):%Y-%m-%d %H:%M}")
wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
try:
    w(f"sheets={wb.sheetnames}")
    all_rows = list(wb["전체당첨내역"].iter_rows(min_row=1, max_col=8, values_only=True))
    w(f"전체당첨내역 rows={len(all_rows)} last3={all_rows[-3:]}")
    for sn in [s for s in wb.sheetnames if s.startswith("3차필터")]:
        ws = wb[sn]
        rows = {r: row_of(ws, r) for r in (1, 2, 3, 4, 6, 7, 8, 9)}
        maxr = ws.max_row
        last = row_of(ws, maxr)
        w(f"\n── {sn}  max_row={maxr} max_col={ws.max_column}")
        w(f"   K2={ws.cell(2, 11).value!r}  K3={ws.cell(3, 11).value!r}  K4={ws.cell(4, 11).value!r}")
        for r in (1, 2, 3, 4, 6, 7, 8, 9):
            row = rows[r]
            if not row:
                w(f"   row{r}: (empty)")
                continue
            lr = row[11:56]
            kinds = {}
            for v in lr:
                k = type(v).__name__
                kinds[k] = kinds.get(k, 0) + 1
            w(f"   row{r}: A={row[0]!r} B..H={row[1:8]} I:K={row[8:11]} L종류={kinds}")
            w(f"           L={short(lr[0], 130)}")
            w(f"           M={short(lr[1], 130)}")
            w(f"           BD={short(lr[44], 130)}")
        w(f"   last row{maxr}: A={last[0]!r} B..H={last[1:8]} I:K={last[8:11]} L0={short(last[11], 80)}")
finally:
    wb.close()

OUT.write_text("\n".join(lines), encoding="utf-8")
print(f"written {OUT}")
