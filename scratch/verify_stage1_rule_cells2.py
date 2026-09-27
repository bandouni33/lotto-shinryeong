# -*- coding: utf-8 -*-
"""1차필터 규칙 추출 v2 — '회차 열에 수식이 있는 행'만 활성 규칙으로 보고,
J가 비어 있으면 수식 안의 문자열 목록에서 targets를 읽는다. 그 뒤 판정을 시트 캐시값과 대조.

v1에서 나온 438건 불일치의 원인:
  · 수식이 없는 행(28, 458, 463~483 …)을 활성 규칙으로 세어 '위반'으로 계산했다 → 제외해야 함
  · 25행처럼 J가 비어 있어도 수식 안에 번호목록이 박혀 있는 행이 있다 → 수식에서 읽어야 함

실행: venv312\\Scripts\\python.exe scratch\\verify_stage1_rule_cells2.py
"""
from __future__ import annotations

import os
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
SRC = ROOT / "★조합생성_후보숫자_추적표" / "조합생성_후보숫자_추적표_샘플.xlsx"
OUT = ROOT / "scratch" / "verify_stage1_rule_cells2_out.txt"
L: list[str] = []
FIRST_ROUND_COL = 14          # N열 (회차 1240)


def w(s: str = "") -> None:
    L.append(s)


def ftext(ws, r, c):
    v = ws.cell(r, c).value
    return v.text if hasattr(v, "text") else v


def targets_from_j(j) -> set[int]:
    return {int(x) for x in re.findall(r"\d+", str(j))} if j is not None else set()


def targets_from_formula(txt: str | None) -> set[int]:
    """수식 안의 따옴표 문자열들 중 쉼표 목록처럼 보이는 것에서 번호를 읽는다."""
    if not txt:
        return set()
    best: set[int] = set()
    for lit in re.findall(r'"([^"]*)"', str(txt)):
        nums = {int(x) for x in re.findall(r"\d+", lit)}
        if len(nums) > len(best):
            best = nums
    return best


wbv = openpyxl.load_workbook(SRC, data_only=True)
wbf = openpyxl.load_workbook(SRC, data_only=False)
try:
    wsa = wbv["전체당첨내역"]
    draws = {}
    for r in range(2, wsa.max_row + 1):
        rr = wsa.cell(r, 1).value
        if isinstance(rr, int):
            draws[rr] = [wsa.cell(r, c).value for c in range(2, 9)]
    wsv = wbv["1차필터(7기본필터)"]
    wsf = wbf["1차필터(7기본필터)"]
    cols = {wsv.cell(4, c).value: c for c in range(1, wsv.max_column + 1)
            if isinstance(wsv.cell(4, c).value, int)}

    rules = []
    inactive = []
    for r in range(5, 1504):
        k, lm = wsv.cell(r, 11).value, wsv.cell(r, 12).value
        if not isinstance(k, (int, float)) or not isinstance(lm, (int, float)):
            continue
        ft = ftext(wsf, r, FIRST_ROUND_COL)
        is_formula = isinstance(ft, str) and ft.startswith("=")
        if not is_formula:
            inactive.append(r)
            continue
        j = wsv.cell(r, 10).value
        tgt = targets_from_j(j) if j is not None else targets_from_formula(ft)
        rules.append({"row": r, "name": wsv.cell(r, 8).value, "j": j,
                      "min": int(k), "max": int(lm), "targets": tgt,
                      "src": "J" if j is not None else "formula"})
    w(f"활성 규칙 {len(rules)}개 (수식 없는 비활성행 {len(inactive)}개: {inactive})")
    w(f"  수식에서 목록을 읽은 행: {[(x['row'], len(x['targets'])) for x in rules if x['src'] == 'formula']}")
    w(f"  J는 있는데 숫자가 0개인 행: {[x['row'] for x in rules if x['src'] == 'J' and not x['targets']]}")

    total = matched = 0
    ex = []
    for rnd, col in sorted(cols.items(), reverse=True):
        nums6 = [draws[rnd][i] for i in range(6)]
        prev7 = draws.get(rnd - 1, [None] * 7)
        for x in rules:
            cached = wsv.cell(x["row"], col).value
            name = str(x["name"] or "")
            if x["row"] == 5:
                tgt = {v for v in prev7 if isinstance(v, int)}
            elif x["row"] in (6, 7):
                tgt = set()
                for v in prev7:
                    if isinstance(v, int):
                        for d in (-1, 0, 1):
                            if 1 <= v + d <= 45:
                                tgt.add(v + d)
            else:
                tgt = x["targets"]
            cnt = len(set(nums6) & tgt)
            viol_mine = not (x["min"] <= cnt <= x["max"])
            viol_sheet = not (cached is None or cached == "")
            total += 1
            if viol_mine == viol_sheet and (not viol_sheet or int(cached) == cnt):
                matched += 1
            elif len(ex) < 15:
                ex.append(f"{rnd}회차 row{x['row']}({name or '-'},src={x['src']}) "
                          f"시트={cached!r} / 내계산 위반={viol_mine} 개수={cnt} "
                          f"(min{x['min']}~max{x['max']}, targets {len(tgt)})")
    w(f"\n검사 {total:,}셀 · 일치 {matched:,} · 불일치 {total - matched:,}")
    for e in ex:
        w("  " + e)
finally:
    wbv.close()
    wbf.close()

OUT.write_text("\n".join(L), encoding="utf-8")
print(f"written {OUT}")
