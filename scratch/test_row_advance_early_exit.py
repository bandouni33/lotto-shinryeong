# -*- coding: utf-8 -*-
"""행 전진 "대상 회차가 없으면 재계산 없이 무시" 불변식 테스트 (합성 워크북 · 빠름).

왜 이 테스트가 필요한가: 진입점 1회 실행이 253초였고 그 절반 이상(약 130초)이 **밀 회차가
하나도 없는데도** 회차 무관 마스크(전체 조합 8,145,060 × 1차 고정 378 × 2차 이격수 48)를
다시 만드는 데 쓰였다. 마스크 생성을 대상 판정 뒤로 옮긴 뒤에도 ①낭비가 실제로 사라졌는지
②순서를 바꾸면서 값·창·오류 처리가 달라지지 않았는지를 여기서 확인한다.

불변식(모든 유효 입력에 성립해야 하는 것):
  I1 대상 없음        → 마스크 생성 0회, 반환 {} (워크북 내용도 무변경)
  I2 대상 있음        → 마스크 생성 **정확히 1회**, 값이 직접 계산한 마스크와 같다
  I3 잘못된 입력 4종  → 예외로 멈추고, **그 전에 마스크를 만들지 않는다**(오류 경로도 싸야 함)
      (a) 추적결과 시트 창 최신이 서로 다름   (b) 회차에 구멍
      (c) 창(102행)보다 많이 밀림             (d) 회차 라벨(1열)이 없음
  I4 values={} 로 advance_row_sheets → 마스크 0회, 라벨 그대로
  I5 전진 뒤 같은 값으로 다시 전진 → 마스크 추가 생성 없음, 재전진 없음(멱등)

실행: venv312\\Scripts\\python.exe scratch\\test_row_advance_early_exit.py
"""
from __future__ import annotations

import itertools
import os
import random
import sys
import time
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

import combo_filter_v2 as cf  # noqa: E402
import weekly_lotto_file_update as wk  # noqa: E402

OUT = ROOT / "scratch" / "test_row_advance_early_exit_out.txt"
MAX_ROUND = 1243
R: list[str] = []
FAILS: list[str] = []
t0 = time.time()


def w(s: str = "") -> None:
    R.append(str(s))
    try:
        print(s, flush=True)
    except UnicodeEncodeError:
        print(str(s).encode("cp949", errors="replace").decode("cp949", errors="replace"),
              flush=True)


def ok(cond: bool, msg: str) -> None:
    if cond:
        w(f"  ok   {msg}")
    else:
        FAILS.append(msg)
        w(f"  FAIL {msg}")


# ── 합성 워크북: 전체당첨내역 40회차 + 행식 추적결과 시트 3개(창 102행)
def make_wb(window_top=None, sheet_tops=None, skip=()) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws_all = wb.active
    ws_all.title = "전체당첨내역"
    r = 2
    for rnd in range(MAX_ROUND - 39, MAX_ROUND + 1):
        if rnd in skip:
            continue
        nums = sorted(random.Random(rnd).sample(range(1, 46), 6))
        ws_all.cell(r, 1, rnd)
        for j, n in enumerate(nums):
            ws_all.cell(r, 2 + j, n)
        ws_all.cell(r, 8, (rnd % 45) + 1)
        r += 1
    top = MAX_ROUND if window_top is None else window_top
    for i, sn in enumerate(wk.ROW_ROUND_SHEETS):
        ws = wb.create_sheet(sn)
        t = top if sheet_tops is None else sheet_tops[i]
        for k in range(wk.ROW_WINDOW):
            ws.cell(wk.ROW_FIRST + k, 1, t - k)
    return wb


_RULES = None


def _rules():
    global _RULES
    if _RULES is None:
        _RULES = cf._load_rules()
    return _RULES


def tiny_base() -> dict:
    """8백만 조합 대신 3,000조합짜리 재료 — 값 정의(규칙 적용·원핫·격차순위)는 그대로다."""
    static_rules, auto_rules, gap_rules = _rules()
    combos = np.array(list(itertools.combinations(range(1, 46), 6))[:3000], dtype=np.int16)
    oh = cf._onehot(combos)
    static_mask, gap_mask = cf.split_static_gap_masks(oh, static_rules, gap_rules, combos)
    return {"combos": combos, "combo_oh": oh, "static_mask": static_mask,
            "gap_mask": gap_mask, "static_rules": static_rules, "auto_rules": auto_rules}


class CountingBase:
    """cf.build_base_masks 를 갈아끼워 '몇 번 만들었는가'를 센다."""

    def __init__(self):
        self.calls = 0
        self.base = tiny_base()

    def __call__(self):
        self.calls += 1
        return self.base


def counted(fn):
    """마스크 캐시를 비우고 카운터를 끼운 채 fn() 실행.
    반환: (결과, 예외, 마스크 생성 횟수, 재료). _tracking_base의 캐시도 비우므로
    '이번 호출이 마스크를 만들었는가'를 정확히 볼 수 있다."""
    counter = CountingBase()
    orig_build, orig_cache = cf.build_base_masks, wk._TRACKING_BASE
    cf.build_base_masks = counter          # type: ignore[assignment]
    wk._TRACKING_BASE = None
    err = out = None
    try:
        out = fn()
    except Exception as e:  # noqa: BLE001
        err = e
    finally:
        cf.build_base_masks = orig_build  # type: ignore[assignment]
        wk._TRACKING_BASE = orig_cache
    return out, err, counter.calls, counter.base


def labels(wb, sn) -> list[int]:
    ws = wb[sn]
    return [ws.cell(r, 1).value for r in range(wk.ROW_FIRST, wk.ROW_LAST + 1)]


w(f"합성 워크북 기준: 전체당첨내역 최신 {MAX_ROUND}회차 · 창 {wk.ROW_WINDOW}행")

# ── I1 대상 없음
w("\n== I1. 대상 회차가 없으면 마스크를 만들지 않는다 ==")
wb = make_wb(window_top=MAX_ROUND)
before = {sn: labels(wb, sn) for sn in wk.ROW_ROUND_SHEETS}
res, err, calls, _ = counted(lambda: wk.row_advance_values(wb, "샘플"))
ok(err is None, f"예외 없음 ({type(err).__name__ if err else '정상'})")
ok(res == {}, f"반환값 = {res} (아무 회차도 계산하지 않는다)")
ok(calls == 0, f"마스크 생성 {calls}회 == 0회 (고치기 전엔 대상이 없어도 1회 만들었다)")
ok({sn: labels(wb, sn) for sn in wk.ROW_ROUND_SHEETS} == before, "워크북 내용 무변경")

# ── I2 대상 있음
w("\n== I2. 대상이 있으면 마스크는 한 번만, 값은 직접 계산과 같다 ==")
wb2 = make_wb(window_top=MAX_ROUND - 1)
vals, err2, calls2, base2 = counted(lambda: wk.row_advance_values(wb2, "샘플"))
ok(err2 is None and sorted(vals) == [MAX_ROUND],
   f"전진 대상 = {sorted(vals)} (창 최신 {MAX_ROUND - 1} 뒤 회차)")
ok(calls2 == 1, f"마스크 생성 {calls2}회 == 1회 (회차 수와 무관하게 재료는 한 번)")
draws = wk._draws_asc_from_sheet(wb2["전체당첨내역"])
ms = cf.compute_stage_masks(draws, MAX_ROUND - 1, base2)
v = vals[MAX_ROUND]
ok(v["stage1"] == int(ms["stage1_mask"].sum()),
   f"1차 통과 {v['stage1']} == 직접 계산 {int(ms['stage1_mask'].sum())}")
ok(v["stage2"] == int(ms["stage2_mask"].sum()),
   f"2차 통과 {v['stage2']} == 직접 계산 {int(ms['stage2_mask'].sum())}")
ok(v["stage4"] == int(ms["stage4_mask"].sum()),
   f"4차 통과 {v['stage4']} == 직접 계산 {int(ms['stage4_mask'].sum())}")
nums = next(h["nums"] for h in draws if h["draw_round"] == MAX_ROUND)
order = ms["gap_order"]
ok(v["K"] == len(set(nums) & set(order[:15])) and v["L"] == len(set(nums) & set(order[15:30]))
   and v["M"] == len(set(nums) & set(order[30:45])),
   f"K/L/M = {v['K']}/{v['L']}/{v['M']} == 격차순위 구간과의 교집합")

# ── I5 멱등(같은 값으로 다시 전진)
w("\n== I5. 전진 뒤 같은 값으로 다시 전진 ==")
wb3 = make_wb(window_top=MAX_ROUND - 1)
vals3, _, _, base3 = counted(lambda: wk.row_advance_values(wb3, "샘플"))
wk._TRACKING_BASE = base3            # 실물 마스크(8백만) 재계산을 피하려고 재료를 캐시에 넣는다
try:
    wk.advance_row_sheets(wb3, "샘플", vals3)
finally:
    wk._TRACKING_BASE = None
top_after = labels(wb3, "1차추적결과")[0]
st2, _, calls_after, _ = counted(lambda: wk.advance_row_sheets(wb3, "샘플", vals3))
ok(top_after == MAX_ROUND, f"전진 결과 창 최신 = {top_after} == {MAX_ROUND}")
ok(all(not s["advanced"] for s in st2.values()), "두 번째 호출은 아무것도 밀지 않는다")
ok(calls_after == 0, f"두 번째 호출의 마스크 생성 {calls_after}회 == 0회(밀 것이 없음)")
ok(labels(wb3, "1차추적결과")[0] == MAX_ROUND, "창 최신 그대로(멱등)")

# ── I4 values={} 경로
w("\n== I4. values={} 경로 ==")
wb4 = make_wb(window_top=MAX_ROUND - 1)
before4 = {sn: labels(wb4, sn) for sn in wk.ROW_ROUND_SHEETS}
st4, _, calls4, _ = counted(lambda: wk.advance_row_sheets(wb4, "샘플", {}))
ok(calls4 == 0, f"마스크 생성 {calls4}회 == 0회")
ok(all(not s["advanced"] for s in st4.values()), "밀지 않음")
ok({sn: labels(wb4, sn) for sn in wk.ROW_ROUND_SHEETS} == before4, "라벨 그대로")

# ── I3 잘못된 입력 4종 (예외 + 마스크 미생성)
w("\n== I3. 잘못된 입력은 예외로 멈추고 마스크도 만들지 않는다 ==")
cases = [
    ("(a) 시트별 창 최신이 다름", make_wb(sheet_tops=(MAX_ROUND, MAX_ROUND, MAX_ROUND - 1))),
    ("(b) 회차에 구멍", make_wb(window_top=MAX_ROUND - 2, skip={MAX_ROUND - 1})),
    ("(c) 창(102행)보다 많이 밀림", make_wb(window_top=MAX_ROUND - wk.ROW_WINDOW - 1)),
]
wb5 = make_wb(window_top=MAX_ROUND)
for sn in wk.ROW_ROUND_SHEETS:
    for r in range(wk.ROW_FIRST, wk.ROW_LAST + 1):
        # 주의: openpyxl의 cell(r, c, None)은 값을 '쓰지 않는다'(기본값이 None이라 생략).
        # 지울 때는 반드시 .value = None 으로 쓴다.
        wb5[sn].cell(r, 1).value = None
cases.append(("(d) 회차 라벨(1열)이 없음", wb5))

for name, bad in cases:
    _, err3, calls3, _ = counted(lambda b=bad: wk.row_advance_values(b, "샘플"))
    ok(err3 is not None, f"{name} → 예외로 멈춤 ({type(err3).__name__ if err3 else '예외 없음'})")
    ok(calls3 == 0, f"{name} → 마스크 생성 {calls3}회 == 0회(오류 경로도 싸다)")

w(f"\n단언 실패 {len(FAILS)}건" + ("" if not FAILS else ": " + " | ".join(FAILS[:5])))
w(f"총 경과 {time.time() - t0:.1f}s")
OUT.write_text("\n".join(R), encoding="utf-8")
print(f"written {OUT}")
os._exit(1 if FAILS else 0)
