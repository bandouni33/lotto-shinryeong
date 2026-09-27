# -*- coding: utf-8 -*-
"""전진 복사본의 '참조 무결성' 정적 검증 (엑셀 없이 — 이 세션에서 엑셀 COM 기동이 안 된다).

엑셀 재계산 없이도 #REF!/#VALUE! 가 날 원인은 정적으로 판정할 수 있다:
  S1 다른 시트 참조가 모두 실존 시트 (원본과 동일한 시트 집합)
  S2 자기 열 참조(N$4·N$1504~N$1516)만 존재하고, 모든 회차 열이 '같은 구조'다
  S3 규칙행 수식이 참조하는 $J/$K/$L(같은 행)에 값이 있다
  S4 모든 회차 라벨과 라벨-1이 전체당첨내역에 존재 = 모든 열이 실제로 계산된다
  S5 구조(열 개수·라벨 연속·수식 셀 수·최대행)가 원본과 동일
  S6 다른 시트 → 컬럼 시트 참조 0건(가장 오래된 열을 버려도 안전)

수식이 수천 자라 같은 문자열을 여러 번 훑으면 느리다 → 회차 열 전체를 '한 번만' 스캔한다.

실행: venv312\\Scripts\\python.exe scratch\\check_sim_advance_static.py
"""
from __future__ import annotations

import re
import sys
import time
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as cl

ROOT = Path(__file__).resolve().parents[1]
# 기본은 시뮬 복사본 비교(검사 대상, 기준). 인자로 두 파일을 바꿔 넣을 수 있다 —
# 실물 적용 뒤에는  실물파일  백업파일  로 돌려 무결성을 재확인한다.
SRC = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else (
    ROOT / "★조합생성_후보숫자_추적표" / "조합생성_후보숫자_추적표_샘플.xlsx")
COPY = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / "scratch" / "_sim_advance_sample.xlsx"
OUT = ROOT / "scratch" / "check_sim_advance_static_out.txt"
TARGET = ("1차필터(7기본필터)", "2차필터(5이격수)")
# 시트명: 따옴표로 감싼 경우와, 앞이 함수명/숫자/기호가 아닌 경우만. '(' ')' 는 넣지 않는다
# (예전 패턴이 'INDEX(전체당첨내역' 같은 가짜 이름을 잡았다 — 검사기 결함).
SHEET_REF = re.compile(r"(?<![A-Za-z0-9_$.])(?:'([^']+)'|([A-Za-z0-9가-힣_]+))!")
COL_ABS_REF = re.compile(r"(?<![A-Za-z0-9_$.!])([A-Z]{1,2})\$(\d+)")

R: list[str] = []
FAILS: list[str] = []
t0 = time.time()


def w(s: str = "") -> None:
    R.append(str(s))
    print(s, flush=True)


def ok(cond: bool, msg: str) -> bool:
    if cond:
        w(f"  ok   {msg}")
    else:
        FAILS.append(msg)
        w(f"  FAIL {msg}")
    return bool(cond)


def fx(cell):
    v = cell.value
    return v.text if hasattr(v, "text") else (v if isinstance(v, str) and v.startswith("=") else None)


def labels_of(ws):
    return [(c, ws.cell(4, c).value) for c in range(1, ws.max_column + 1)
            if isinstance(ws.cell(4, c).value, int)]


def scan_sheet(ws):
    """회차 열 범위를 한 번만 훑어 참조 통계를 모은다."""
    lab = labels_of(ws)
    first, last = lab[0][0], lab[-1][0]
    cross: Counter = Counter()          # 다른 시트 참조
    own: dict[str, Counter] = {}        # 열별 자기 열 참조(행번호)
    other_cols: Counter = Counter()     # 자기 열이 아닌 열 참조(있으면 번역 누락)
    n_formula = 0
    for c in range(first, last + 1):
        own_letter = cl(c)
        cnt: Counter = Counter()
        for r in range(1, ws.max_row + 1):
            t = fx(ws.cell(r, c))
            if not t:
                continue
            n_formula += 1
            for m in SHEET_REF.finditer(t):
                cross[m.group(1) or m.group(2)] += 1
            for letter, row in COL_ABS_REF.findall(t):
                if letter == own_letter:
                    cnt[int(row)] += 1
                else:
                    other_cols[letter] += 1
        own[own_letter] = cnt
    return {"first": first, "last": last, "labels": [v for _, v in lab],
            "cross": cross, "own": own, "other_cols": other_cols, "n_formula": n_formula}


wbc = openpyxl.load_workbook(COPY, data_only=False)
wbs = openpyxl.load_workbook(SRC, data_only=False)
try:
    names = wbc.sheetnames
    wsa = wbc["전체당첨내역"]
    rounds = {int(wsa.cell(r, 1).value) for r in range(2, wsa.max_row + 1)
              if isinstance(wsa.cell(r, 1).value, int)}
    w(f"검사 대상 {COPY.name} · 기준 {SRC.name}")
    w(f"복사본 시트 {len(names)}개 · 전체당첨내역 회차 {min(rounds)}~{max(rounds)}")

    # S6 다른 시트 → 컬럼 시트 참조 0건
    refs_in = []
    for sn in names:
        if sn in TARGET:
            continue
        for row in wbc[sn].iter_rows():
            for c in row:
                t = fx(c)
                if t and any(f"{n}!" in t for n in TARGET):
                    refs_in.append(f"{sn}!{c.coordinate}")
    total = 1
    ok(not refs_in, f"S6 다른 시트에서 컬럼 시트로의 참조 0건 (발견 {len(refs_in)} {refs_in[:3]})")

    for sn in TARGET:
        ws, ws_src = wbc[sn], wbs[sn]
        st, st_src = scan_sheet(ws), scan_sheet(ws_src)
        w(f"\n── {sn}: 회차열 {cl(st['first'])}..{cl(st['last'])} "
          f"({len(st['labels'])}개) · 수식 {st['n_formula']:,}개 · max_row {ws.max_row} "
          f"({time.time() - t0:.0f}s)")

        total += 1
        ok(set(st["cross"]) == set(st_src["cross"]) and set(st["cross"]) <= set(names),
           f"S1 다른 시트 참조 = {sorted(st['cross'])} (원본과 동일, 모두 실존 시트)")

        total += 1
        ok(not st["other_cols"],
           f"S2 자기 열이 아닌 '열 고정' 참조 0건 (발견 {dict(st['other_cols'])})")

        # 모든 회차 열이 같은 자기 열 참조 구조 + 범위 안
        shapes = Counter(tuple(sorted(c.items())) for c in st["own"].values())
        maxrow_ref = max((r for c in st["own"].values() for r in c), default=0)
        total += 2
        ok(len(shapes) == 1,
           f"S2 회차 {len(st['own'])}개 열의 자기 열 참조 구조가 모두 동일 "
           f"(서로 다른 구조 {len(shapes)}종)")
        ok(maxrow_ref <= ws.max_row,
           f"S2 자기 열 참조 최대 행 {maxrow_ref} <= 시트 최대행 {ws.max_row}")

        total += 1
        miss = [v for v in st["labels"] if v not in rounds or (v - 1) not in rounds]
        ok(not miss, f"S4 모든 라벨/직전 회차가 전체당첨내역에 존재 (빠짐 {miss[:5]})")

        if sn.startswith("1차필터"):
            bad_jkl = [f"{nm}{r}" for r in range(5, 1504)
                       if fx(ws.cell(r, st["first"]))
                       for col, nm in ((10, "J"), (11, "K"), (12, "L"))
                       if ws.cell(r, col).value is None]
            total += 1
            ok(not bad_jkl, f"S3 규칙행 수식이 참조하는 J/K/L 값 존재 (빠짐 {len(bad_jkl)} {bad_jkl[:5]})")

        total += 4
        ok(len(st["labels"]) == len(st_src["labels"]),
           f"S5 회차 열 개수 원본과 동일 ({len(st['labels'])} == {len(st_src['labels'])})")
        ok(st["labels"] == [st["labels"][0] - i for i in range(len(st["labels"]))],
           f"S5 라벨 1씩 연속 감소 ({st['labels'][0]}..{st['labels'][-1]})")
        ok(st["n_formula"] == st_src["n_formula"],
           f"S5 수식 셀 수 원본과 동일 ({st['n_formula']:,} == {st_src['n_formula']:,})")
        ok(ws.max_row == ws_src.max_row, f"S5 최대행 불변 ({ws.max_row} == {ws_src.max_row})")
finally:
    wbc.close()
    wbs.close()

w(f"\n단언 {total}건 · 실패 {len(FAILS)}건" + ("" if not FAILS else ": " + " | ".join(FAILS[:5])))
w(f"총 경과 {time.time() - t0:.0f}s")
OUT.write_text("\n".join(R), encoding="utf-8")
print(f"written {OUT}")
sys.exit(1 if FAILS else 0)
