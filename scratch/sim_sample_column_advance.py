# -*- coding: utf-8 -*-
"""샘플 파일 '회차 컬럼 전진' 불변식 테스트 — 임시 복사본에서만 실행(실물 무변경).

전진 로직은 생산 코드 한 곳(weekly_lotto_file_update.shift_round_columns /
advance_round_columns)에만 있다. 이 스크립트는 그 **실제 자동화 함수**를 임시
복사본에 돌려 불변식을 단언한다(사본 로직을 따로 두지 않는다 — 검증 대상이
자동화 코드와 어긋나면 의미가 없기 때문).

전진 1회 = 새 회차 1개를 컬럼으로 추가하고 가장 오래된 컬럼을 버린다:
  · 새 회차 열 = 옛 최신 열 원문 그대로(자기 열 참조라 번역 불필요)
  · 밀리는 열은 자기 열 참조만 +1 번역(N$4→O$4)  · 가장 오래된 열은 비운다
  · 서식·열 너비·병합·숨김 열은 함께 옮긴다

단언: 열 개수/라벨 연속, 수식 셀 행 집합 동일, 번역 문자 단위 일치, 새 N열 == 옛 N열,
헬퍼행 1504~1516, 병합·너비·서식·숨김 열, 옛 열 참조 잔존 없음, 다른 시트 무변경,
최종 라벨 == 전체당첨내역 최신, 오류값 0, 실물 파일 무변경.

실행: venv312\\Scripts\\python.exe scratch\\sim_sample_column_advance.py
결과 복사본: scratch/_sim_advance_sample.xlsx
"""
from __future__ import annotations

import hashlib
import re
import shutil
import sys
import time
from copy import copy as _style_copy
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as cl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

# 생산 코드의 전진 로직을 그대로 쓴다(단일 기준점).
from weekly_lotto_file_update import (  # noqa: E402
    COLUMN_ROUND_SHEETS,
    _cell_formula as cell_formula,
    _round_columns as round_columns,
    _translate_col_refs as translate_text,
    shift_round_columns,
)


class _CellLike:
    """생산 코드의 _cell_formula는 '셀 객체'를 받는다 — 값만 들고 있는 대역."""

    __slots__ = ("value",)

    def __init__(self, value):
        self.value = value


def formula_text(value):
    """셀 값 → 수식 문자열(None 허용). 추출 로직은 생산 코드 것을 그대로 쓴다."""
    return cell_formula(_CellLike(value)) if value is not None else None

SRC = ROOT / "★조합생성_후보숫자_추적표" / "조합생성_후보숫자_추적표_샘플.xlsx"
OUT_COPY = ROOT / "scratch" / "_sim_advance_sample.xlsx"
OUT_TXT = ROOT / "scratch" / "sim_sample_column_advance_out.txt"

R: list[str] = []
FAILS: list[str] = []
t0 = time.time()


def w(s: str = "") -> None:
    R.append(str(s))
    print(s, flush=True)


def ok(cond: bool, msg: str) -> bool:
    if cond:
        w(f"  ok   {msg}")
    else:
        FAILS.append(msg)
        w(f"  FAIL {msg}")
    return bool(cond)


def snapshot_column(ws, col: int, max_row: int):
    items = []
    for r in range(1, max_row + 1):
        cell = ws.cell(r, col)
        if cell.value is not None:
            items.append((r, cell.value, _style_copy(cell._style)))
    return items


def sheet_cell_hash(wb, name: str) -> str:
    h = hashlib.sha256()
    for row in wb[name].iter_rows():
        for c in row:
            v = c.value
            t = v.text if hasattr(v, "text") else v
            h.update(f"{c.coordinate}={t!r}|".encode())
    return h.hexdigest()


# ══════════════════════════ 진단 모드 ══════════════════════════
if "--probe-merges" in sys.argv:
    wb0 = openpyxl.load_workbook(SRC, data_only=False)
    try:
        for sn in COLUMN_ROUND_SHEETS:
            ws0 = wb0[sn]
            f0, l0, lb0 = round_columns(ws0)
            w(f"── {sn}: 회차열 {cl(f0)}..{cl(l0)} ({len(lb0)}개) 라벨 {lb0[0]}→{lb0[-1]}")
            for rng in ws0.merged_cells.ranges:
                if rng.min_col <= l0 and rng.max_col >= f0:
                    w(f"   merged {rng.coord}: 값={ws0.cell(rng.min_row, rng.min_col).value!r}")
            hidden = [c for c in range(1, ws0.max_column + 1)
                      if ws0.column_dimensions[cl(c)].hidden]
            w(f"   숨긴 열: {[cl(c) for c in hidden]}")
    finally:
        wb0.close()
    sys.exit(0)

# ══════════════════════════ 실행 ══════════════════════════
SRC_SHA = hashlib.sha256(SRC.read_bytes()).hexdigest()
shutil.copy2(SRC, OUT_COPY)
w(f"실물 sha(앞12) = {SRC_SHA[:12]} · 복사본 {OUT_COPY.name}")
w("전진 로직: weekly_lotto_file_update.shift_round_columns (생산 코드 직접 호출)")

wb = openpyxl.load_workbook(OUT_COPY, data_only=False)
try:
    sheet_names_before = list(wb.sheetnames)
    keep_hashes = {sn: sheet_cell_hash(wb, sn) for sn in wb.sheetnames
                   if sn not in COLUMN_ROUND_SHEETS}
    max_row = max(wb[sn].max_row for sn in COLUMN_ROUND_SHEETS)
    wsa = wb["전체당첨내역"]
    latest = max(int(wsa.cell(r, 1).value) for r in range(2, wsa.max_row + 1)
                 if isinstance(wsa.cell(r, 1).value, int))
    w(f"시트 {len(sheet_names_before)}개 · 최신 회차 {latest}")

    total = 0
    for ws_name in COLUMN_ROUND_SHEETS:
        ws = wb[ws_name]
        first, last, labels = round_columns(ws)
        merges_before = sorted(rng.coord for rng in ws.merged_cells.ranges)
        hidden_before = [c for c in range(1, ws.max_column + 1)
                         if ws.column_dimensions[cl(c)].hidden]
        base = labels[0]
        n_shift = latest - base
        w(f"\n── {ws_name}: 회차열 {cl(first)}..{cl(last)} ({len(labels)}개) "
          f"라벨 {labels[0]}→{labels[-1]} · 수식 {sum(1 for c in range(first, last + 1) for r in range(1, max_row + 1) if formula_text(ws.cell(r, c).value)):,}개 "
          f"· 전체당첨내역 최신 {latest}까지 {n_shift}회 밀기")

        for step in range(1, n_shift + 1):
            target = base + step
            # 밀기 '직전' 상태를 스냅샷(검증 기준)
            before_text = {c: {r: formula_text(ws.cell(r, c).value) for r in range(1, max_row + 1)
                               if formula_text(ws.cell(r, c).value)}
                           for c in range(first, last + 1)}
            before_style = {c: {r: ws.cell(r, c)._style for r in range(1, max_row + 1)
                                if formula_text(ws.cell(r, c).value)}
                            for c in range(first, last + 1)}
            before_rows = {c: set(t) for c, t in before_text.items()}
            before_widths = {c: round(ws.column_dimensions[cl(c)].width or 0, 4)
                             for c in range(first, last + 1)}

            st = shift_round_columns(ws, target)          # ← 생산 코드
            f_, l_ = st["first"], st["last"]
            w(f"  [{ws_name}] {step}/{n_shift}회 밀기 → N4={target}")

            new_labels = [ws.cell(4, c).value for c in range(f_, l_ + 1)]
            total += 1
            ok(len(new_labels) == len(labels) and ws.cell(4, l_ + 2).value is None,
               f"{ws_name}: 회차 열 개수 {len(new_labels)} == {len(labels)} (범위 {cl(f_)}..{cl(l_)} 불변)")
            total += 1
            ok(new_labels == [target - i for i in range(len(new_labels))],
               f"{ws_name}: 라벨 연속 {new_labels[0]}..{new_labels[-1]} (1씩 감소)")

            bad = []
            for c in range(f_, l_ + 1):
                src_c = c if c == f_ else c - 1
                now = {r for r in range(1, max_row + 1) if formula_text(ws.cell(r, c).value)}
                if now != before_rows[src_c]:
                    bad.append((cl(c), sorted(before_rows[src_c] - now)[:3], sorted(now - before_rows[src_c])[:3]))
            total += 1
            ok(not bad, f"{ws_name}: 수식 셀 행 집합 전 열 동일 (불일치 {len(bad)}열 {bad[:2]})")

            mism = []
            for c in range(f_ + 1, l_ + 1):
                src_letter, dst_letter = cl(c - 1), cl(c)
                for r, old_text in before_text[c - 1].items():
                    want = translate_text(old_text, src_letter, dst_letter)
                    got = formula_text(ws.cell(r, c).value)
                    if got != want:
                        mism.append((f"{dst_letter}{r}", want[:50], str(got)[:50]))
            total += 1
            ok(not mism, f"{ws_name}: 밀린 열 수식 = '자기 열만 +1' 번역과 문자 단위 일치 (불일치 {len(mism)} {mism[:1]})")

            n_mism = [r for r, old_text in before_text[f_].items()
                      if formula_text(ws.cell(r, f_).value) != old_text]
            total += 1
            ok(not n_mism, f"{ws_name}: 새 {cl(f_)}열 수식 == 옛 {cl(f_)}열 원문 (불일치 {len(n_mism)} {n_mism[:3]})")

            if ws_name.startswith("1차필터"):
                present = [r for r in range(1504, 1517) if formula_text(ws.cell(r, f_).value)]
                refs_self = [r for r in present
                             if re.search(r"(?<![A-Za-z0-9_$.!])" + cl(f_) + r"\$4",
                                          formula_text(ws.cell(r, f_).value))]
                total += 2
                ok(len(present) == 13, f"{ws_name}: 헬퍼행 1504~1516 수식 13개 존재 (실제 {len(present)})")
                ok(len(refs_self) == len(present),
                   f"{ws_name}: 헬퍼행이 모두 {cl(f_)}$4(자기 열 라벨)를 참조")

            left = []
            for c in range(f_ + 1, l_ + 1):
                src_letter = cl(c - 1)
                for r in range(1, max_row + 1):
                    t = formula_text(ws.cell(r, c).value)
                    if t and re.search(r"(?<![A-Za-z0-9_$.!])" + src_letter + r"\$?\d", t):
                        left.append(f"{cl(c)}{r}←{src_letter}")
            total += 1
            ok(not left, f"{ws_name}: 밀린 열에 옛 열 참조 잔존 없음 (발견 {len(left)} {left[:2]})")

            st_mism = [f"{cl(c)}{r}" for c in range(f_, l_ + 1)
                       for r in before_style[c if c == f_ else c - 1]
                       if ws.cell(r, c)._style != before_style[c if c == f_ else c - 1][r]]
            total += 1
            ok(not st_mism, f"{ws_name}: 수식 셀 서식 원본과 동일 (불일치 {len(st_mism)} {st_mism[:3]})")

            total += 1
            ok(sorted(rng.coord for rng in ws.merged_cells.ranges) == merges_before,
               f"{ws_name}: 병합 범위 불변 {merges_before}")
            total += 1
            w_mism = [(cl(c), before_widths[c if c == f_ else c - 1],
                       round(ws.column_dimensions[cl(c)].width or 0, 4))
                      for c in range(f_, l_ + 1)
                      if round(ws.column_dimensions[cl(c)].width or 0, 4)
                      != before_widths[c if c == f_ else c - 1]]
            ok(not w_mism, f"{ws_name}: 열 너비가 내용과 함께 옮겨짐 (불일치 {len(w_mism)} {w_mism[:2]})")
            total += 1
            ok([c for c in range(1, ws.max_column + 1) if ws.column_dimensions[cl(c)].hidden]
               == hidden_before, f"{ws_name}: 숨김 열 목록 불변")

            labels = new_labels

    changed = [sn for sn in sheet_names_before if sn not in COLUMN_ROUND_SHEETS
               and sheet_cell_hash(wb, sn) != keep_hashes[sn]]
    total += 2
    ok(not changed, f"회차 컬럼 시트 외 내용 무변경 (변경 {changed})")
    ok(list(wb.sheetnames) == sheet_names_before, "시트 목록 불변")

    last_labels = [wb[sn].cell(4, round_columns(wb[sn])[0]).value for sn in COLUMN_ROUND_SHEETS]
    total += 1
    ok(all(v == latest for v in last_labels),
       f"밀기 후 최신 라벨 {last_labels} == 전체당첨내역 최신 회차 {latest}")

    errs = [f"{sn}!{c.coordinate}" for sn in wb.sheetnames for row in wb[sn].iter_rows()
            for c in row if isinstance(c.value, str) and c.value.strip() in (
                "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!", "#N/A")]
    total += 1
    ok(not errs, f"오류값 0 (발견 {len(errs)} {errs[:5]})")

    w(f"\n단언 {total}건 · 실패 {len(FAILS)}건")
    wb.save(OUT_COPY)
finally:
    wb.close()

total += 1
ok(hashlib.sha256(SRC.read_bytes()).hexdigest() == SRC_SHA, "실물 파일 무변경(복사본에서만 작업)")
w(f"복사본 저장: {OUT_COPY} ({OUT_COPY.stat().st_size:,} bytes) · 총 경과 {time.time() - t0:.1f}s")
w(f"실패 {len(FAILS)}건" + ("" if not FAILS else ": " + " | ".join(FAILS[:5])))

OUT_TXT.write_text("\n".join(R), encoding="utf-8")
print(f"written {OUT_TXT}")
sys.exit(1 if FAILS else 0)
