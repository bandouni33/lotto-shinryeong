# -*- coding: utf-8 -*-
"""대상 4+1 파일의 '행 단위 회차 시트' 유무·창 크기 실측 (읽기 전용).

행 전진(1차추적결과/2차추적결과/4차필터)을 어느 파일에 붙여야 하는지, 창 크기와
최신 라벨이 파일마다 같은지 확인한다. 실물 파일은 열지 않고 읽기만 한다.
"""
from __future__ import annotations

import hashlib
import os
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)

DIR = ROOT / "★조합생성_후보숫자_추적표"
FILES = [
    DIR / "조합생성_후보숫자_추적표_샘플.xlsx",
    DIR / "조합생성_후보숫자_추적표_샘플_200회검증용.xlsx",
    DIR / "조합생성_후보숫자_추적표_전체표본_윈도우비교.xlsx",
    DIR / "조합생성_후보숫자_추적표_최근500표본_윈도우비교.xlsx",
    DIR / "★후보숫자_추적표_표본vs최근50회.xlsx",
]
ROW_SHEETS = ("1차추적결과", "2차추적결과", "4차필터")
OUT = ROOT / "scratch" / "probe_row_sheets_files_out.txt"
R: list[str] = []


def w(s: str = "") -> None:
    R.append(str(s))
    print(s, flush=True)


for p in FILES:
    if not p.exists():
        w(f"== {p.name}: 없음")
        continue
    w(f"== {p.name}  sha={hashlib.sha256(p.read_bytes()).hexdigest()[:12]}")
    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    try:
        sns = list(wb.sheetnames)
        w(f"   시트 {len(sns)}개: {sns}")
        wsa = wb["전체당첨내역"] if "전체당첨내역" in sns else None
        last = None
        if wsa is not None:
            rounds = [r0 for r0 in
                      (ws_row[0] for ws_row in wsa.iter_rows(min_row=2, max_col=1, values_only=True))
                      if isinstance(r0, int)]
            last = max(rounds) if rounds else None
        w(f"   전체당첨내역 최신 = {last}")
        for sn in ROW_SHEETS:
            if sn not in sns:
                continue
            ws = wb[sn]
            labels = [ws.cell(r, 1).value for r in range(3, min(ws.max_row, 120) + 1)]
            ints = [x for x in labels if isinstance(x, int)]
            w(f"   [{sn}] max_row={ws.max_row} max_col={ws.max_column} "
              f"A3={ws.cell(3, 1).value} 최신라벨={ints[0] if ints else None} "
              f"최오래라벨={ints[-1] if ints else None} 라벨수={len(ints)}")
        for sn in sns:
            if sn.startswith("3차필터") or sn.startswith("기준"):
                ws = wb[sn]
                w(f"   [{sn}] max_row={ws.max_row} A1={ws.cell(1, 1).value} "
                  f"K2={ws.cell(2, 11).value} K3={ws.cell(3, 11).value}")
    finally:
        wb.close()

OUT.write_text("\n".join(R), encoding="utf-8")
print(f"written {OUT}")
