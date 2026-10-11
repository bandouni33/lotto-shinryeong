"""1-4차 필터링 표본 파일 주간 자동 갱신(2026-10-11 사용자 지시).

  W1 시트 고르기: L:BD 표준(3차/4차필터) ↔ S:BK 보드 ↔ 행 시트(없는 '4차필터'는 건너뜀)
  W2 보드 시트 따라잡기: 밀린 회차 확정·다음 대기 행·창 크기 유지·J 수식·오른쪽 패널/병합 그대로·멱등
  W3 외부 파일 참조('[1]전체당첨내역!') → 자기 파일 참조, 다 고치면 외부 연결 제거
  W4 FILE_PATHS 에 표본 파일 등록, 옛 200회검증용 빠짐

실행: python tests/test_weekly_pyobon.py
"""

from __future__ import annotations

import random
import sys
import unittest
from pathlib import Path

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import weekly_lotto_file_update as w  # noqa: E402

w.log = lambda *_a, **_k: None


def _draws(n=260, seed=7):
    rnd = random.Random(seed)
    return [{"draw_round": r, "nums": sorted(rnd.sample(range(1, 46), 6)), "bonus": 1}
            for r in range(1, n + 1)]


def _book(draws, pending_round, rows=10):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws_all = wb.create_sheet("전체당첨내역")
    ws_all.append(["회차", 1, 2, 3, 4, 5, 6, "보너스"])
    for h in draws:
        ws_all.append([h["draw_round"], *h["nums"], h["bonus"]])
    b = wb.create_sheet("3차필터(200출현빈도순 후보)")
    for i in range(45):
        b.cell(1, 19 + i).value = i + 1
    b.cell(3, 19).value = ("=SUMPRODUCT((전체당첨내역!$A$2:$A$1500>=(MAX(전체당첨내역!$A$2:$A$1500)-199))"
                           "*(전체당첨내역!$B$2:$B$1500=S$1))")
    b.cell(6, 1).value = "회차"
    b.cell(7, 1).value = pending_round
    for k in range(1, rows):
        b.cell(7 + k, 1).value = pending_round - k
        b.cell(7 + k, 8).value = 999            # 옛 H값(고정)
    b["BO16"] = "패널"
    b.merge_cells("BO16:BP19")
    s = wb.create_sheet("4차필터(100vs30회_후보)")
    s.cell(6, 1).value = "회차"
    s.cell(1, 12).value, s.cell(1, 56).value = 1, 45
    s.cell(7, 12).value = ArrayFormula(ref="L7", text="=MAX([1]전체당첨내역!$A$2:$A$1500)")
    s.cell(2, 12).value = "=COUNTIF([1]전체당첨내역!$B$2:$G$1500,L$1)"
    wb.create_sheet("1차추적결과")
    return wb


class PyobonTest(unittest.TestCase):
    def test_W1_sheet_selection(self):
        wb = _book(_draws(), 261)
        self.assertEqual(w.std_forecast_sheets(wb), ["4차필터(100vs30회_후보)"])
        self.assertEqual(w.board_forecast_sheets(wb), ["3차필터(200출현빈도순 후보)"])
        self.assertEqual(w.row_sheets_in_wb(wb, "t"), ["1차추적결과"])

    def test_W2_board_catch_up(self):
        draws = _draws()
        wb = _book(draws, 258)                   # 258·259·260 밀림
        ws = wb["3차필터(200출현빈도순 후보)"]
        self.assertEqual(w._board_recent_window(ws), 200)
        done = w.catch_up_board_sheet(ws, draws)
        self.assertEqual(done, [258, 259, 260])
        self.assertEqual([ws.cell(r, 1).value for r in range(7, 17)], list(range(261, 251, -1)))
        self.assertIsNone(ws.cell(17, 1).value)  # 창 크기 유지(10행)
        self.assertEqual([ws.cell(7, c).value for c in range(19, 64)], w.gap_forecast(draws, 260, 200))
        board = [ws.cell(8, c).value for c in range(19, 64)]
        self.assertEqual(board, w.gap_forecast(draws, 259, 200))
        got = set(draws[259]["nums"])
        self.assertEqual(ws.cell(8, 11).value, len(got & set(board[:15])))
        self.assertEqual(ws.cell(8, 10).value, "=SUM(K8:M8)")
        self.assertEqual(ws.cell(12, 10).value, "=SUM(K12:M12)")
        self.assertEqual(ws.cell(11, 8).value, 999)   # 지난 회차 값 고정
        self.assertIsNone(ws.cell(8, 8).value)        # 새 행 H 는 비움
        self.assertEqual(ws["BO16"].value, "패널")
        self.assertIn("BO16:BP19", [str(r) for r in ws.merged_cells.ranges])
        self.assertEqual(w.catch_up_board_sheet(ws, draws), [])   # 멱등

    def test_W3_localize_external_refs(self):
        wb = _book(_draws(), 261)
        wb._external_links = [object()]
        self.assertEqual(w.localize_external_refs(wb), 2)
        s = wb["4차필터(100vs30회_후보)"]
        self.assertNotIn("[1]", s.cell(7, 12).value.text)
        self.assertNotIn("[1]", s.cell(2, 12).value)
        self.assertEqual(wb._external_links, [])
        self.assertEqual(w.localize_external_refs(wb), 0)

    def test_W4_registered(self):
        self.assertIn("1-4차표본", w.FILE_PATHS)
        self.assertEqual(w.FILE_PATHS["1-4차표본"].name, "★1-4차 필터링 표본.xlsx")
        self.assertNotIn("200회검증용", w.FILE_PATHS)
        self.assertIn("1-4차표본", w.ADVANCE_3CHA_FILES)
        self.assertIn("1-4차표본", w.ROW_ROUND_FILES)
        self.assertNotIn("1-4차표본", w.DAEJO_FILES)


if __name__ == "__main__":
    unittest.main(verbosity=2)
